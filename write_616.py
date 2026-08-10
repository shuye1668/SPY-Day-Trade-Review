import openpyxl
from datetime import datetime

# === trades_all.xlsx ===
wb = openpyxl.load_workbook(r'C:\TradeReview\trades_all.xlsx')
ws = wb['Sheet1']

trades_616 = [
    ('2026-06-16','10:07:41','SPY',754.5865,None,None,100,'B','0','2026-06-16','10:26:34'),
    ('2026-06-16','10:18:13','SPY',754.41,None,None,100,'B','0','2026-06-16','10:19:48'),
    ('2026-06-16','10:19:48','SPY',753.93,'多','-49.57',100,'S','0','2026-06-16','10:18:13'),
    ('2026-06-16','10:26:34','SPY',753.8601,'多','-74.21',100,'S','0','2026-06-16','10:07:41'),
    ('2026-06-16','10:48:16','SPY',753.13,None,None,100,'B','1',None,None),
    ('2026-06-16','10:59:38','SPY',752.755,None,None,100,'B','0','2026-06-16','11:00:35'),
    ('2026-06-16','11:00:35','SPY',753.127,'多','+35.63',100,'S','0','2026-06-16','10:59:38'),
    ('2026-06-16','11:28:42','SPY',752.535,None,None,100,'B','0','2026-06-16','11:29:24'),
    ('2026-06-16','11:29:24','SPY',752.825,'多','+27.43',100,'S','0','2026-06-16','11:28:42'),
    ('2026-06-16','11:46:12','SPY',752.285,None,None,100,'B','0','2026-06-16','11:49:20'),
    ('2026-06-16','11:49:20','SPY',752.17,'多','-13.07',100,'S','0','2026-06-16','11:46:12'),
    ('2026-06-16','11:51:27','SPY',752.33,None,None,100,'B','0','2026-06-16','11:51:59'),
    ('2026-06-16','11:51:59','SPY',752.6601,'多','+31.44',100,'S','0','2026-06-16','11:51:27'),
]

start_row = 714
for i, t in enumerate(trades_616):
    r = start_row + i
    ws.cell(row=r, column=1, value=t[0])
    ws.cell(row=r, column=2, value=t[1])
    ws.cell(row=r, column=3, value=t[2])
    ws.cell(row=r, column=4, value=t[3])
    if t[4] is not None:
        ws.cell(row=r, column=5, value=t[4])
    if t[5] is not None:
        ws.cell(row=r, column=6, value=t[5])
    ws.cell(row=r, column=7, value=t[6])
    ws.cell(row=r, column=8, value=t[7])
    ws.cell(row=r, column=9, value=t[8])
    if t[9] is not None:
        ws.cell(row=r, column=10, value=t[9])
    if t[10] is not None:
        ws.cell(row=r, column=11, value=t[10])

wb.save(r'C:\TradeReview\trades_all.xlsx')
print('trades_all.xlsx written: rows 714-726')

# === CS交易紀錄.xlsx - 個別識別法 sheet ===
wb2 = openpyxl.load_workbook(r'C:\TradeReview\CS交易紀錄.xlsx')
ws2 = wb2['個別識別法']

# Day Header at row 1418 (2 blank rows after last content row 1415)
ws2.cell(row=1418, column=1, value='2026.06.16')
ws2.cell(row=1418, column=2, value=-115970.06)
ws2.cell(row=1418, column=3, value='損益')

# Detail rows 1419-1431
detail_data = [
    (-75458.65, None),
    (-75441.00, None),
    (75391.43, -49.57),
    (75384.44, -74.21),
    (-75313.00, None),
    (-75275.50, None),
    (75311.13, 35.63),
    (-75253.50, None),
    (75280.93, 27.43),
    (-75228.50, None),
    (75215.43, -13.07),
    (-75233.00, None),
    (75264.44, 31.44),
]

for i, (col1, col3) in enumerate(detail_data):
    r = 1419 + i
    ws2.cell(row=r, column=1, value=col1)
    if r == 1419:
        ws2.cell(row=r, column=2, value='=+B1418+A1419')
    else:
        ws2.cell(row=r, column=2, value=f'=+B{r-1}+A{r}')
    if col3 is not None:
        ws2.cell(row=r, column=3, value=col3)

# 合計 row 1432
ws2.cell(row=1432, column=2, value='合計 ')
ws2.cell(row=1432, column=3, value='=SUM(C1419:C1431)')

# 核驗 row 1433
ws2.cell(row=1433, column=2, value='核驗')
ws2.cell(row=1433, column=3, value='=+B1431-B1418')

# 當沖 summary rows 1435-1438
# Row 1435: 買 (SUMIF minus hold entry at row 1423)
ws2.cell(row=1435, column=1, value='2026.06.16當沖')
ws2.cell(row=1435, column=2, value='買')
ws2.cell(row=1435, column=3, value='=SUMIF(A1419:A1431,"<0")-A1423')

# Row 1436: 賣
ws2.cell(row=1436, column=1, value='="SPY " & COUNT(C1419:C1431) * 100 & "股"')
ws2.cell(row=1436, column=2, value='賣')
ws2.cell(row=1436, column=3, value='=SUMIF(A1419:A1431,">0")')

# Row 1437: 當日損益
ws2.cell(row=1437, column=2, value='當日損益')
ws2.cell(row=1437, column=3, value='=+C1435+C1436')

# Row 1438: 當日損益 %
ws2.cell(row=1438, column=2, value='當日損益 %')
ws2.cell(row=1438, column=3, value='=C1437/ABS(C1435)')

# 留倉 summary rows 1440-1441
ws2.cell(row=1440, column=1, value='累積留倉   100股')
ws2.cell(row=1441, column=1, value='2026.06.16 留倉')
ws2.cell(row=1441, column=2, value='買 100股')
ws2.cell(row=1441, column=3, value=-75313.00)

wb2.save(r'C:\TradeReview\CS交易紀錄.xlsx')
print('CS交易紀錄.xlsx 個別識別法 written: rows 1418-1441')

# === CS交易紀錄.xlsx - 累積損益check sheet ===
wb3 = openpyxl.load_workbook(r'C:\TradeReview\CS交易紀錄.xlsx')
ws3 = wb3['累積損益check']

ws3.cell(row=28, column=1, value=datetime(2026, 6, 16))
ws3.cell(row=28, column=2, value=-42.35)
ws3.cell(row=28, column=3, value='=C27+B28')

wb3.save(r'C:\TradeReview\CS交易紀錄.xlsx')
print('CS交易紀錄.xlsx 累積損益check written: row 28')

print('ALL WRITES COMPLETE')
