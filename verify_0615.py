# -*- coding: utf-8 -*-
import openpyxl

print("=== trades_all.xlsx verification ===")
wb = openpyxl.load_workbook('trades_all.xlsx')
ws = wb.active
print(f"max_row: {ws.max_row}")

print("\nUpdated hold rows (Status should be '2'):")
for r in [616, 617, 636, 689]:
    print(f"  Row {r}: Status={ws.cell(r,9).value}, Pair_Date={ws.cell(r,10).value}, Pair_Time={ws.cell(r,11).value}")

print("\nNew rows 706-713:")
for r in range(706, 714):
    vals = [ws.cell(r, c).value for c in range(1, 12)]
    print(f"  Row {r}: {vals}")

# Check no remaining Status='1'
holds = []
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 9).value == '1':
        holds.append(r)
print(f"\nRemaining holds (Status='1'): {holds if holds else 'None (stack empty)'}")

print("\n=== CS交易紀錄.xlsx verification ===")
wb2 = openpyxl.load_workbook('CS交易紀錄.xlsx')
ws2 = wb2['個別識別法']
print(f"max_row: {ws2.max_row}")

print("\n6/15 block:")
for r in range(1380, ws2.max_row + 1):
    vals = [ws2.cell(r, c).value for c in range(1, 5)]
    if any(v is not None for v in vals):
        print(f"  Row {r}: {vals}")

ws3 = wb2['累積損益check']
print(f"\n累積損益check last row ({ws3.max_row}):")
print(f"  {[ws3.cell(ws3.max_row, c).value for c in range(1, 5)]}")
