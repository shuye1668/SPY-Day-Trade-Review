import openpyxl
wb = openpyxl.load_workbook('CS交易紀錄.xlsx')
ws = wb['個別識別法']
print('max_row:', ws.max_row)
for r in range(1351, 1379):
    print(f'Row {r}:', [ws.cell(r, c).value for c in range(1, 5)])
