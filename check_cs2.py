# -*- coding: utf-8 -*-
import openpyxl
wb = openpyxl.load_workbook('CS交易紀錄.xlsx')
ws = wb['個別識別法']

# Check the actual cell values for 留倉 references
# Row 1375 C = '=C1346', Row 1376 C = '=C1347', Row 1377 C = '=C1348', Row 1378 C = '=C1349'
# Need to check what rows 1346-1349 contain (these are from 6/11 block's detail rows)
print("6/11 detail rows referenced by 6/12 留倉:")
for r in range(1344, 1352):
    print(f'Row {r}:', [ws.cell(r, c).value for c in range(1, 5)])

# Also check 累積損益check sheet
ws2 = wb['累積損益check']
print("\n累積損益check last rows:")
print('max_row:', ws2.max_row)
for r in range(max(1, ws2.max_row - 5), ws2.max_row + 1):
    print(f'Row {r}:', [ws2.cell(r, c).value for c in range(1, 6)])
