#!/usr/bin/env python3
"""把 CS交易紀錄.xlsx 匯出成純文字側寫，讓 git diff 看得懂帳本改了什麼。

xlsx 是 zip 容器，git 只會顯示 "Bin 12345 -> 12350 bytes"。帳本每天都在動、
而且會被人工修正，看不到「哪一格從 X 變成 Y」等於失去版本控制最大的價值。

這支腳本產生 CS交易紀錄_dump.txt：每列一行、公式與值都保留，純文字可 diff。
它是唯讀操作，永遠不會動到 xlsx 本身。

用法：
    python dump_cs_text.py                 # 產出 CS交易紀錄_dump.txt
    python dump_cs_text.py --check         # 只檢查是否與現有 dump 不同（exit 1=有異動）
"""
import argparse
import os
import sys

import openpyxl

_BASE = os.path.dirname(os.path.abspath(__file__))
CS_PATH = os.path.join(_BASE, "CS交易紀錄.xlsx")
OUT_PATH = os.path.join(_BASE, "CS交易紀錄_dump.txt")


def cell_repr(c):
    v = c.value
    if v is None:
        return ""
    if isinstance(v, float):
        # 統一小數位，避免浮點尾差造成整份 dump 假異動
        return f"{v:.6f}".rstrip("0").rstrip(".")
    return str(v)


def dump(cs_path=CS_PATH):
    # data_only=False：保留公式字串（=+B123+A124），那才是帳本結構的本體。
    # 取值版本會隨 Excel 是否重算而變，反而不穩定。
    wb = openpyxl.load_workbook(cs_path, data_only=False)
    lines = []
    for name in wb.sheetnames:
        ws = wb[name]
        last = 0
        for r in range(ws.max_row, 0, -1):
            if any(ws.cell(r, c).value not in (None, "") for c in range(1, 6)):
                last = r
                break
        lines.append(f"### SHEET {name}  (rows 1..{last})")
        for r in range(1, last + 1):
            cells = [cell_repr(ws.cell(r, c)) for c in range(1, 6)]
            if not any(cells):
                lines.append(f"{r}\t")           # 保留空行位置，行號才對得上
            else:
                lines.append(f"{r}\t" + "\t".join(cells))
        lines.append("")
    return "\n".join(lines) + "\n"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cs", default=CS_PATH)
    ap.add_argument("--out", default=OUT_PATH)
    ap.add_argument("--check", action="store_true",
                    help="只比對不寫入；有異動回 exit 1")
    a = ap.parse_args()
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    if not os.path.exists(a.cs):
        print(f"找不到 {a.cs}")
        return 1

    text = dump(a.cs)
    old = ""
    if os.path.exists(a.out):
        with open(a.out, encoding="utf-8") as f:
            old = f.read()

    if a.check:
        changed = (old != text)
        print("CHANGED" if changed else "UNCHANGED")
        return 1 if changed else 0

    with open(a.out, "w", encoding="utf-8", newline="\n") as f:
        f.write(text)
    n = text.count("\n")
    print(f"已寫出 {os.path.basename(a.out)}：{n} 行"
          + ("（內容有異動）" if old != text else "（無異動）"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
