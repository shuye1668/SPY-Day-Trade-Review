# -*- coding: utf-8 -*-
"""Rebuild CS交易紀錄 blocks from 5/27 onwards: ALL TRDs in detail, SUMIF adjusted."""

import openpyxl
import shutil
from collections import defaultdict

src = r'C:\TradeReview\CS交易紀錄.xlsx'
shutil.copy2(src, src.replace('.xlsx', '_backup_rebuild.xlsx'))

# === Read trades_all ===
wb_trades = openpyxl.load_workbook(r'C:\TradeReview\trades_all.xlsx')
ws_trades = wb_trades.active

trades = []
for r in range(2, ws_trades.max_row + 1):
    date_val = ws_trades.cell(r, 1).value
    if date_val is None:
        continue
    d = str(date_val)
    if d < '2026-05-27':
        continue
    pnl_raw = ws_trades.cell(r, 6).value
    trade = {
        'date': d,
        'time': str(ws_trades.cell(r, 2).value),
        'price': float(ws_trades.cell(r, 4).value),
        'pnl_str': str(pnl_raw).strip() if pnl_raw is not None else None,
        'action': str(ws_trades.cell(r, 8).value),
        'status': str(ws_trades.cell(r, 9).value),
        'pair_date': str(ws_trades.cell(r, 10).value) if ws_trades.cell(r, 10).value else None,
        'pair_time': str(ws_trades.cell(r, 11).value) if ws_trades.cell(r, 11).value else None,
    }
    trades.append(trade)

trade_lookup = {}
for t in trades:
    trade_lookup[(t['date'], t['time'])] = t

trades_by_date = defaultdict(list)
for t in trades:
    trades_by_date[t['date']].append(t)
for d in trades_by_date:
    trades_by_date[d].sort(key=lambda x: x['time'])

b_header = {
    '2026-05-27': -42252.81,
    '2026-05-28': -42213.54,
    '2026-05-29': -42112.31,
    '2026-06-03': -42019.94,
    '2026-06-04': -41859.07,
    '2026-06-05': -41673.22,
    '2026-06-08': -41589.50,
}

days_order = ['2026-05-27', '2026-05-28', '2026-05-29',
              '2026-06-03', '2026-06-04', '2026-06-05', '2026-06-08']

# === Open CS file and delete old blocks ===
wb = openpyxl.load_workbook(src)
ws = wb['個別識別法']

max_r = ws.max_row
if max_r >= 1067:
    ws.delete_rows(1067, max_r - 1067 + 1)

cur = 1067
hold_rows = {}       # (date, time) -> row in 留倉 section
active_holds = []    # trades currently held overnight

for day in days_order:
    day_trades = trades_by_date.get(day, [])
    if not day_trades:
        continue

    ds = day.replace('-', '.')

    hold_entries = [t for t in day_trades if t['action'] == 'B' and t['status'] in ('1', '2')]
    close_exits  = [t for t in day_trades if t['action'] == 'S' and t['status'] == '2']

    # --- compute detail ---
    detail = []
    for t in day_trades:
        if t['action'] == 'B':
            c1 = round(-(t['price'] * 100), 2)
            pnl = None
        else:
            paired = trade_lookup.get((t['pair_date'], t['pair_time']))
            ep = paired['price'] if paired else 0
            pv = float(t['pnl_str'].replace('+', '')) if t['pnl_str'] else 0
            c1 = round(pv + ep * 100, 2)
            pnl = pv if t['status'] == '0' else None
        detail.append({'t': t, 'c1': c1, 'pnl': pnl})

    # --- Day Header ---
    ws.cell(cur, 1, ds);  ws.cell(cur, 2, b_header[day]);  ws.cell(cur, 3, '損益')
    hdr = cur;  cur += 1

    # --- Detail ---
    fd = cur
    he_rows = []   # hold entry rows in detail
    ce_rows = []   # close exit rows in detail
    for dd in detail:
        ws.cell(cur, 1, dd['c1'])
        ws.cell(cur, 2, f'=+B{cur-1}+A{cur}')
        if dd['pnl'] is not None:
            ws.cell(cur, 3, dd['pnl'])
        if dd['t']['action'] == 'B' and dd['t']['status'] in ('1', '2'):
            he_rows.append(cur)
        if dd['t']['action'] == 'S' and dd['t']['status'] == '2':
            ce_rows.append(cur)
        cur += 1
    ld = cur - 1

    # --- 合計 / 核驗 ---
    ws.cell(cur, 2, '合計 ');  ws.cell(cur, 3, f'=SUM(C{fd}:C{ld})');  cur += 1
    ws.cell(cur, 2, '核驗');   ws.cell(cur, 3, f'=+B{ld}-B{hdr}');     cur += 1
    cur += 1  # blank

    # --- 當沖 ---
    br = cur
    ws.cell(cur, 1, f'{ds}當沖');  ws.cell(cur, 2, '買')
    sb = f'=SUMIF(A{fd}:A{ld},"<0")'
    for h in he_rows:
        sb += f'-A{h}'
    ws.cell(cur, 3, sb);  cur += 1

    sr = cur
    ws.cell(cur, 1, f'="SPY " & COUNT(C{fd}:C{ld}) * 100 & "股"');  ws.cell(cur, 2, '賣')
    ss = f'=SUMIF(A{fd}:A{ld},">0")'
    for c in ce_rows:
        ss += f'-A{c}'
    ws.cell(cur, 3, ss);  cur += 1

    ws.cell(cur, 2, '當沖損益');   ws.cell(cur, 3, f'=+C{br}+C{sr}');  pr = cur;  cur += 1
    ws.cell(cur, 2, '當沖損益 %'); ws.cell(cur, 3, f'=C{pr}/ABS(C{br})');         cur += 1
    cur += 1  # blank after 當沖

    # --- 平倉 ---
    if close_exits:
        for ce in close_exits:
            paired = trade_lookup.get((ce['pair_date'], ce['pair_time']))
            if not paired:
                continue
            eds = paired['date'].replace('-', '.')
            hk = (paired['date'], paired['time'])

            ws.cell(cur, 1, f'{eds}買入');  ws.cell(cur, 2, '100股')
            if hk in hold_rows:
                ws.cell(cur, 3, f'=C{hold_rows[hk]}')
            else:
                ws.cell(cur, 3, round(-(paired['price'] * 100), 2))
            er = cur;  cur += 1

            pv = float(ce['pnl_str'].replace('+', ''))
            ec1 = round(pv + paired['price'] * 100, 2)
            ws.cell(cur, 1, f'{ds}平倉');  ws.cell(cur, 2, '100股');  ws.cell(cur, 3, ec1)
            xr = cur;  cur += 1

            ws.cell(cur, 2, '平倉損益');   ws.cell(cur, 3, f'=+C{er}+C{xr}');  cpr = cur;  cur += 1
            ws.cell(cur, 2, '平倉損益 %'); ws.cell(cur, 3, f'=C{cpr}/ABS(C{er})');          cur += 1
        cur += 1  # blank after 平倉

    # --- update active holds ---
    closed = {(ce['pair_date'], ce['pair_time']) for ce in close_exits}
    active_holds = [h for h in active_holds if (h['date'], h['time']) not in closed]
    for he in hold_entries:
        active_holds.append(he)

    # --- 留倉 ---
    if active_holds:
        ws.cell(cur, 1, f'累積留倉   {len(active_holds)*100}股');  cur += 1
        for h in active_holds:
            hds = h['date'].replace('-', '.')
            hk = (h['date'], h['time'])
            ws.cell(cur, 1, f'{hds} 留倉');  ws.cell(cur, 2, f'買 100股')
            if hk in hold_rows:
                ws.cell(cur, 3, f'=C{hold_rows[hk]}')
            else:
                ws.cell(cur, 3, round(-(h['price'] * 100), 2))
            hold_rows[hk] = cur
            cur += 1
        cur += 1  # blank after 留倉

    cur += 1  # second trailing blank

wb.save(src)
print(f'Rebuilt rows 1067-{cur-3}. Active holds: {len(active_holds)}')

# === Verify ===
wb2 = openpyxl.load_workbook(src, data_only=True)
ws2 = wb2['個別識別法']
for r in range(1067, cur):
    vals = [ws2.cell(r, c).value for c in range(1, 4)]
    if any(v is not None for v in vals):
        print(f'  {r}: {vals}')
