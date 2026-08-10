import openpyxl
from datetime import datetime, timedelta
import sys
sys.stdout.reconfigure(encoding='utf-8')

###############################################################################
# 1. VERIFICATION — qty * price ~ abs(amount), tolerance $0.02
###############################################################################
trades_raw = [
    # (date_csv, time_tw, action, qty, symbol, price, misc_fees, amount, ref)
    ('6/4/26', '22:14:11', 'BOT', 100, 'SPY', 753.175,   0,    -75317.50, '1006606840770'),
    ('6/4/26', '22:23:39', 'SOLD', 100, 'SPY', 753.86,   1.57,  75386.00, '1006607291593'),
    ('6/4/26', '22:29:30', 'BOT', 100, 'SPY', 754.0544,  0,    -75405.44, '1006607743204'),
    ('6/4/26', '22:52:51', 'SOLD', 100, 'SPY', 753.7901,  1.57,  75379.01, '1006608921970'),
    ('6/4/26', '23:06:12', 'BOT', 100, 'SPY', 754.7288,  0,    -75472.88, '1006609793178'),
    ('6/4/26', '23:07:09', 'SOLD', 100, 'SPY', 754.81,   1.57,  75481.00, '1006609793427'),
    ('6/4/26', '23:10:15', 'BOT',  60, 'SPY', 755.17,    0,    -45310.20, '1006610240223'),
    ('6/4/26', '23:10:15', 'BOT',  40, 'SPY', 755.17,    0,    -30206.80, '1006610240223'),
    ('6/4/26', '23:11:02', 'SOLD', 100, 'SPY', 755.4101,  1.58,  75541.01, '1006610240406'),
    ('6/4/26', '23:13:42', 'BOT', 100, 'SPY', 755.459,   0,    -75545.90, '1006610241023'),
    ('6/4/26', '23:15:18', 'SOLD', 100, 'SPY', 755.61,   1.58,  75561.00, '1006610241459'),
    ('6/5/26', '00:05:34', 'BOT', 100, 'SPY', 755.54,    0,    -75554.00, '1006612633335'),
    ('6/5/26', '00:23:42', 'SOLD', 100, 'SPY', 756.60,   1.58,  75660.00, '1006613109913'),
]

print("=== VERIFICATION ===")
all_pass = True
for t in trades_raw:
    date, time, action, qty, sym, price, misc, amount, ref = t
    expected = qty * price
    actual = abs(amount)
    diff = abs(expected - actual)
    if diff > 0.02:
        all_pass = False
        print(f"  FAIL: {action} {qty} {sym} @{price} | expected={expected:.2f} actual={actual:.2f} diff={diff:.2f}")

if all_pass:
    print(f"  All {len(trades_raw)} TRD rows pass (tolerance $0.02)")
else:
    print("  STOPPING: verification failed!")
    sys.exit(1)

###############################################################################
# 2. MERGE SPLIT FILLS + EDT CONVERSION
###############################################################################
def tw_to_edt(time_str, date_str):
    h, m, s = map(int, time_str.split(':'))
    total_min = (h * 60 + m - 12 * 60) % (24 * 60)
    edt_h, edt_m = divmod(total_min, 60)
    parts = date_str.split('/')
    month, day, year = int(parts[0]), int(parts[1]), 2000 + int(parts[2])
    if h < 12:
        d = datetime(year, month, day) - timedelta(days=1)
        us_date = d.strftime('%Y-%m-%d')
    else:
        us_date = f"{year}-{month:02d}-{day:02d}"
    edt_time = f"{edt_h:02d}:{edt_m:02d}:{s:02d}"
    return edt_time, us_date

merged_trades = []
i = 0
while i < len(trades_raw):
    t = trades_raw[i]
    date, time, action, qty, sym, price, misc, amount, ref = t
    total_qty = qty
    total_amount = amount
    total_misc = misc
    j = i + 1
    while j < len(trades_raw) and trades_raw[j][8] == ref:
        total_qty += trades_raw[j][3]
        total_amount += trades_raw[j][7]
        total_misc += trades_raw[j][6]
        j += 1
    if j > i + 1:
        avg_price = abs(total_amount) / total_qty
        print(f"  Split fill merged: Ref#{ref}, {j-i} fills, {total_qty} shares @{avg_price:.4f}")
    else:
        avg_price = price
    edt_time, us_date = tw_to_edt(time, date)
    col1 = total_amount - abs(total_misc)
    merged_trades.append({
        'us_date': us_date,
        'edt_time': edt_time,
        'symbol': sym,
        'price': avg_price if j > i + 1 else price,
        'qty': total_qty,
        'action': 'B' if action == 'BOT' else 'S',
        'col1': col1,
        'amount': total_amount,
        'misc_fees': total_misc,
    })
    i = j

print(f"\nMerged: {len(merged_trades)} trades (from {len(trades_raw)} raw lines)")
print(f"US trading day: {merged_trades[0]['us_date']}")

###############################################################################
# 3. LIFO MATCHING
###############################################################################
print("\n=== LIFO MATCHING ===")
stack = []
pairs = []

for idx, t in enumerate(merged_trades):
    if not stack:
        stack.append(idx)
    elif (t['action'] == 'S' and merged_trades[stack[-1]]['action'] == 'B') or \
         (t['action'] == 'B' and merged_trades[stack[-1]]['action'] == 'S'):
        entry_idx = stack.pop()
        entry = merged_trades[entry_idx]
        pnl = t['col1'] + entry['col1']
        type_label = '多' if entry['action'] == 'B' else '空'
        pairs.append({
            'entry_idx': entry_idx, 'exit_idx': idx,
            'pnl': round(pnl, 2), 'type': type_label, 'status': 0,
        })
        merged_trades[entry_idx]['pair_idx'] = idx
        merged_trades[idx]['pair_idx'] = entry_idx
        merged_trades[idx]['pnl'] = round(pnl, 2)
        merged_trades[idx]['type'] = type_label
        merged_trades[entry_idx]['type'] = None
        merged_trades[entry_idx]['pnl'] = None
        merged_trades[idx]['status'] = 0
        merged_trades[entry_idx]['status'] = 0
    else:
        stack.append(idx)

if stack:
    print(f"  WARNING: {len(stack)} overnight holds remain")
    for idx in stack:
        merged_trades[idx]['status'] = 1
        merged_trades[idx]['type'] = None
        merged_trades[idx]['pnl'] = None
else:
    print("  All positions closed. No overnight holds.")

total_pnl = 0
winners = 0
losers = 0
for p in pairs:
    entry = merged_trades[p['entry_idx']]
    exit_t = merged_trades[p['exit_idx']]
    pnl = p['pnl']
    total_pnl += pnl
    if pnl > 0: winners += 1
    else: losers += 1
    sign = '+' if pnl >= 0 else ''
    print(f"  {entry['action']}@{entry['price']:.4f} -> {exit_t['action']}@{exit_t['price']:.4f} = {sign}{pnl:.2f} ({p['type']})")

print(f"\n  Total PnL: {'+' if total_pnl >= 0 else ''}{total_pnl:.2f}")
print(f"  Winners: {winners}, Losers: {losers}")
wr = winners / (winners + losers) * 100
print(f"  Win rate: {winners}/{winners+losers} = {wr:.1f}%")

###############################################################################
# 4. WRITE trades_all.xlsx
###############################################################################
print("\n=== WRITING trades_all.xlsx ===")
wb = openpyxl.load_workbook(r'C:\TradeReview\trades_all.xlsx')
ws = wb.active
start_row = ws.max_row + 1
print(f"  Appending {len(merged_trades)} rows starting at row {start_row}")

for i, t in enumerate(merged_trades):
    r = start_row + i
    partner = merged_trades[t.get('pair_idx', i)]
    pnl_str = None
    if t.get('pnl') is not None:
        pnl_val = t['pnl']
        pnl_str = f"+{pnl_val:.2f}" if pnl_val >= 0 else f"{pnl_val:.2f}"
    ws.cell(r, 1, t['us_date'])
    ws.cell(r, 2, t['edt_time'])
    ws.cell(r, 3, t['symbol'])
    ws.cell(r, 4, t['price'])
    ws.cell(r, 5, t.get('type'))
    ws.cell(r, 6, pnl_str)
    ws.cell(r, 7, t['qty'])
    ws.cell(r, 8, t['action'])
    ws.cell(r, 9, str(t['status']))
    ws.cell(r, 10, partner['us_date'])
    ws.cell(r, 11, partner['edt_time'])

wb.save(r'C:\TradeReview\trades_all.xlsx')
print("  trades_all.xlsx saved.")

###############################################################################
# 5. WRITE CS交易紀錄.xlsx
###############################################################################
print("\n=== WRITING CS交易紀錄.xlsx ===")
wb2 = openpyxl.load_workbook(r'C:\TradeReview\CS交易紀錄.xlsx')
ws2 = wb2['個別識別法']

target_header = '2026.06.04'
for r in range(1, ws2.max_row + 1):
    if ws2.cell(r, 1).value == target_header and ws2.cell(r, 3).value == '損益':
        print(f"  Already contains {target_header} at row {r}. Skipping.")
        wb2.close()
        sys.exit(0)

b_0604 = -41859.07
last_row = ws2.max_row
header_row = last_row + 3

print(f"  Header row: {header_row}, B={b_0604}")

ws2.cell(header_row, 1, target_header)
ws2.cell(header_row, 2, b_0604)
ws2.cell(header_row, 3, '損益')

detail_start = header_row + 1
for i, t in enumerate(merged_trades):
    r = detail_start + i
    ws2.cell(r, 1, round(t['col1'], 2))
    if i == 0:
        ws2.cell(r, 2, f'=+B{header_row}+A{r}')
    else:
        ws2.cell(r, 2, f'=+B{r-1}+A{r}')
    if t.get('pnl') is not None:
        ws2.cell(r, 3, round(t['pnl'], 2))

detail_end = detail_start + len(merged_trades) - 1

r_total = detail_end + 1
ws2.cell(r_total, 2, '合計 ')
ws2.cell(r_total, 3, f'=SUM(C{detail_start}:C{detail_end})')

r_verify = r_total + 1
ws2.cell(r_verify, 2, '核驗')
ws2.cell(r_verify, 3, f'=+B{detail_end}-B{header_row}')

r_blank1 = r_verify + 1

r_buy = r_blank1 + 1
ws2.cell(r_buy, 1, '2026.06.04當沖')
ws2.cell(r_buy, 2, '買')
ws2.cell(r_buy, 3, f'=SUMIF(A{detail_start}:A{detail_end},"<0")')

r_sell = r_buy + 1
ws2.cell(r_sell, 1, f'="SPY " & COUNT(C{detail_start}:C{detail_end}) * 100 & "股"')
ws2.cell(r_sell, 2, '賣')
ws2.cell(r_sell, 3, f'=SUMIF(A{detail_start}:A{detail_end},">0")')

r_pnl = r_sell + 1
ws2.cell(r_pnl, 2, '當日損益')
ws2.cell(r_pnl, 3, f'=+C{r_buy}+C{r_sell}')

r_pct = r_pnl + 1
ws2.cell(r_pct, 2, '當日損益 %')
ws2.cell(r_pct, 3, f'=C{r_pnl}/ABS(C{r_buy})')

print(f"  Detail: rows {detail_start}-{detail_end}")
print(f"  合計: row {r_total}, 核驗: row {r_verify}")
print(f"  SPY 當沖: rows {r_buy}-{r_pct}")

ws3 = wb2['累積損益check']
next_row = ws3.max_row + 1
ws3.cell(next_row, 1, datetime(2026, 6, 4))
ws3.cell(next_row, 2, round(total_pnl, 2))
ws3.cell(next_row, 3, f'=C{next_row-1}+B{next_row}')
print(f"  累積損益check: row {next_row} added")

wb2.save(r'C:\TradeReview\CS交易紀錄.xlsx')
print("  CS交易紀錄.xlsx saved.")

print(f"\n=== DONE ===")
print(f"US Trading Day: 2026-06-04 (Thursday)")
print(f"Trades: {len(merged_trades)} ({len(pairs)} round trips)")
print(f"Total PnL: {'+' if total_pnl >= 0 else ''}{total_pnl:.2f}")
print(f"Win rate: {winners}/{winners+losers} = {wr:.1f}%")
