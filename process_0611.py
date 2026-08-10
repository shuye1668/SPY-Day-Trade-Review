import openpyxl
from collections import OrderedDict

# Raw TRD data for US trading day 6/11 (lines 80-100 from CSV)
# (taiwan_date, taiwan_time, action, qty, symbol, price, ref, amount, misc_fees, balance)
raw_trds = [
    ('6/11/26','21:36:33','BOT',40,'SPY',727.785,'1006707339215',-29111.40,0,-145767.03),
    ('6/11/26','21:36:33','BOT',60,'SPY',727.785,'1006707339215',-43667.10,0,-189434.13),
    ('6/11/26','21:49:23','SOLD',100,'SPY',729.21,'1006708221649',72921.00,1.52,-116514.65),
    ('6/11/26','21:57:29','BOT',100,'SPY',730.51,'1006709188325',-73051.00,0,-189565.65),
    ('6/11/26','22:01:18','SOLD',100,'SPY',730.92,'1006709189462',73092.00,1.53,-116475.18),
    ('6/11/26','22:20:01','BOT',100,'SPY',727.03,'1006710687208',-72703.00,0,-189178.18),
    ('6/11/26','22:22:58','SOLD',100,'SPY',728.55,'1006711110179',72855.00,1.52,-116324.70),
    ('6/11/26','22:30:13','BOT',100,'SPY',729.09,'1006711484638',-72909.00,0,-189233.70),
    ('6/11/26','22:35:39','SOLD',100,'SPY',728.67,'1006711944038',72867.00,1.52,-116368.22),
    ('6/11/26','22:37:06','BOT',100,'SPY',729.859,'1006711944509',-72985.90,0,-189354.12),
    ('6/11/26','22:42:57','SOLD',100,'SPY',729.6709,'1006712328434',72967.09,1.52,-116388.55),
    ('6/11/26','22:56:03','BOT',100,'SPY',726.4348,'1006713016700',-72643.48,0,-189032.03),
    ('6/11/26','23:07:55','SOLD',40,'SPY',726.93,'1006713830277',29077.20,0.61,-159955.44),
    ('6/11/26','23:07:55','SOLD',60,'SPY',726.92,'1006713830277',43615.20,0.91,-116341.15),
    ('6/11/26','23:10:51','BOT',100,'SPY',727.0299,'1006713831062',-72702.99,0,-189044.14),
    ('6/11/26','23:11:50','SOLD',100,'SPY',727.115,'1006713831307',72711.50,1.52,-116334.16),
    ('6/11/26','23:12:27','BOT',100,'SPY',727.0799,'1006713831527',-72707.99,0,-189042.15),
    ('6/11/26','23:13:08','SOLD',100,'SPY',727.51,'1006713831700',72751.00,1.52,-116292.67),
    ('6/12/26','00:33:06','BOT',100,'SPY',727.9459,'1006717854773',-72794.59,0,-189087.26),
    ('6/12/26','00:43:51','BOT',100,'SPY',728.535,'1006718253151',-72853.50,0,-261940.76),
    ('6/12/26','01:29:03','SOLD',100,'SPY',729.83,'1006720250147',72983.00,1.52,-188959.28),
]

# Verify qty * price ~ abs(amount)
print('=== Verification ===')
all_pass = True
for t in raw_trds:
    date,time,action,qty,sym,price,ref,amount,fees,bal = t
    expected = qty * price
    actual = abs(amount)
    diff = abs(expected - actual)
    if diff > 0.02:
        print(f'FAIL: {date} {time} {action} {qty}@{price}: expected={expected:.2f} actual={actual:.2f} diff={diff:.4f}')
        all_pass = False
print(f'All {len(raw_trds)} rows pass verification: {all_pass}')

# Merge split fills by Ref#
merged = OrderedDict()
for t in raw_trds:
    date,time,action,qty,sym,price,ref,amount,fees,bal = t
    key = ref
    if key not in merged:
        merged[key] = {'date':date,'time':time,'action':action,'qty':0,'sym':sym,
                       'total_amount':0,'total_fees':0,'ref':ref,'last_bal':bal}
    m = merged[key]
    m['qty'] += qty
    m['total_amount'] += amount
    m['total_fees'] += fees
    m['last_bal'] = bal

trades = []
for ref, m in merged.items():
    avg_price = abs(m['total_amount']) / m['qty']
    col1 = m['total_amount'] - abs(m['total_fees'])
    tw_h, tw_m, tw_s = map(int, m['time'].split(':'))
    total_min = (tw_h * 60 + tw_m - 12 * 60) % (24 * 60)
    edt_h, edt_m = divmod(total_min, 60)
    edt_time = f'{edt_h:02d}:{edt_m:02d}:{tw_s:02d}'

    trades.append({
        'date': '2026-06-11',
        'edt_time': edt_time,
        'action': 'B' if m['action'] == 'BOT' else 'S',
        'price': round(avg_price, 4),
        'qty': m['qty'],
        'col1': col1,
        'amount': m['total_amount'],
        'fees': m['total_fees'],
        'ref': ref,
    })

print(f'\n=== Merged trades: {len(trades)} ===')
for t in trades:
    print(f"  {t['edt_time']} {t['action']} {t['qty']}@{t['price']:.4f} col1={t['col1']:.2f}")

# LIFO pairing
stack = []
stack.append({'date':'2026-06-05','edt_time':'10:44:16','action':'B','price':749.37,'qty':100,'is_prior':True,'xlsx_row':616,'col1':-74937.00})
stack.append({'date':'2026-06-05','edt_time':'11:14:45','action':'B','price':748.45,'qty':100,'is_prior':True,'xlsx_row':617,'col1':-74845.00})
stack.append({'date':'2026-06-09','edt_time':'10:12:37','action':'B','price':742.915,'qty':100,'is_prior':True,'xlsx_row':636,'col1':-74291.50})

pairs = []
holds = []

for t in trades:
    if t['action'] == 'B':
        stack.append({'date':t['date'],'edt_time':t['edt_time'],'action':'B','price':t['price'],
                       'qty':t['qty'],'col1':t['col1'],'is_prior':False,'trade':t})
    else:
        if stack and stack[-1]['action'] == 'B':
            entry = stack.pop()
            pnl = round(t['col1'] + entry['col1'], 2)
            pair_type = '平' if entry['is_prior'] else '多'
            status = '2' if entry['is_prior'] else '0'
            pairs.append({
                'entry': entry,
                'exit': t,
                'pnl': pnl,
                'type': pair_type,
                'status': status,
            })
        else:
            print(f"WARNING: SOLD with no matching BOT at {t['edt_time']}")

for s in stack:
    if not s['is_prior']:
        holds.append(s)

print(f'\n=== LIFO Results ===')
print(f'Day trade pairs: {len([p for p in pairs if p["status"]=="0"])}')
print(f'Close prior hold: {len([p for p in pairs if p["status"]=="2"])}')
print(f'New overnight holds: {len(holds)}')
print(f'Remaining prior holds: {len([s for s in stack if s["is_prior"]])}')

total_dt_pnl = sum(p['pnl'] for p in pairs if p['status'] == '0')
wins = len([p for p in pairs if p['status'] == '0' and p['pnl'] > 0])
losses = len([p for p in pairs if p['status'] == '0' and p['pnl'] < 0])
n_dt = len([p for p in pairs if p['status'] == '0'])
print(f'\nDay trade PnL: {total_dt_pnl:+.2f}')
print(f'Win/Loss: {wins}/{losses} ({wins/n_dt*100:.1f}%)')

for p in pairs:
    e = p['entry']
    x = p['exit']
    print(f"  {p['type']} B@{e['price']:.4f} -> S@{x['price']:.4f} PnL={p['pnl']:+.2f}")

for h in holds:
    print(f"  Hold: B@{h['price']:.4f} at {h['date']} {h['edt_time']}")

# ============================================================
# STEP 2: Write to trades_all.xlsx
# ============================================================
print('\n=== Writing to trades_all.xlsx ===')
wb = openpyxl.load_workbook('trades_all.xlsx')
ws = wb.active
start_row = ws.max_row + 1
print(f'Starting at row {start_row}')

# Build rows in chronological order
xlsx_rows = []

# Collect all events with their EDT times for sorting
events = []
for p in pairs:
    e = p['entry']
    x = p['exit']
    if not e['is_prior']:
        events.append(('entry', e['edt_time'], e, p))
    events.append(('exit', x['edt_time'], x, p))

for h in holds:
    events.append(('hold', h['edt_time'], h, None))

events.sort(key=lambda ev: ev[1])

written_entries = set()
row_num = start_row

for ev_type, edt_time, data, pair in events:
    if ev_type == 'entry':
        entry = data
        exit_trade = pair['exit']
        ws.cell(row_num, 1, entry['date'])
        ws.cell(row_num, 2, entry['edt_time'])
        ws.cell(row_num, 3, 'SPY')
        ws.cell(row_num, 4, entry['price'])
        ws.cell(row_num, 5, None)
        ws.cell(row_num, 6, None)
        ws.cell(row_num, 7, 100)
        ws.cell(row_num, 8, 'B')
        ws.cell(row_num, 9, pair['status'])
        ws.cell(row_num, 10, exit_trade['date'] if isinstance(exit_trade, dict) and 'date' in exit_trade else '2026-06-11')
        ws.cell(row_num, 11, exit_trade['edt_time'] if isinstance(exit_trade, dict) else exit_trade)
        print(f'  Row {row_num}: B {entry["edt_time"]} @{entry["price"]:.4f} Status={pair["status"]}')
        row_num += 1
    elif ev_type == 'exit':
        exit_t = data
        entry = pair['entry']
        ws.cell(row_num, 1, '2026-06-11')
        ws.cell(row_num, 2, exit_t['edt_time'])
        ws.cell(row_num, 3, 'SPY')
        ws.cell(row_num, 4, exit_t['price'])
        ws.cell(row_num, 5, pair['type'])
        pnl_str = f"{pair['pnl']:+.2f}" if pair['pnl'] != 0 else '0.00'
        ws.cell(row_num, 6, pnl_str)
        ws.cell(row_num, 7, 100)
        ws.cell(row_num, 8, 'S')
        ws.cell(row_num, 9, pair['status'])
        entry_time = entry['edt_time']
        entry_date = entry['date']
        ws.cell(row_num, 10, entry_date)
        ws.cell(row_num, 11, entry_time)
        print(f'  Row {row_num}: S {exit_t["edt_time"]} @{exit_t["price"]:.4f} {pair["type"]} PnL={pnl_str} Status={pair["status"]}')
        row_num += 1
    elif ev_type == 'hold':
        h = data
        ws.cell(row_num, 1, h['date'])
        ws.cell(row_num, 2, h['edt_time'])
        ws.cell(row_num, 3, 'SPY')
        ws.cell(row_num, 4, h['price'])
        ws.cell(row_num, 5, None)
        ws.cell(row_num, 6, None)
        ws.cell(row_num, 7, 100)
        ws.cell(row_num, 8, 'B')
        ws.cell(row_num, 9, '1')
        ws.cell(row_num, 10, None)
        ws.cell(row_num, 11, None)
        print(f'  Row {row_num}: B {h["edt_time"]} @{h["price"]:.4f} HOLD')
        row_num += 1

wb.save('trades_all.xlsx')
print(f'trades_all.xlsx saved. Total rows: {row_num - 1}')
