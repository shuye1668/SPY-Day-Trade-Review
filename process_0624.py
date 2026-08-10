import openpyxl
from datetime import datetime

DATE_STR = '2026-06-24'
DATE_DOT = '2026.06.24'
DAY_HEADER_B = 36912.80
MISC_FEE_SPY = 1.54

# ============================================================
# 1. trades_all.xlsx - Append 20 SPY rows for 6/24
# ============================================================
pairs = [
    ('空', 'S', 735.8106, '09:46:52', 'B', 734.87,   '09:48:20',  92.52),
    ('多', 'B', 734.51,   '09:49:04', 'S', 734.9101, '09:49:37',  38.47),
    ('多', 'B', 736.105,  '09:50:24', 'S', 736.5601, '09:50:45',  43.97),
    ('空', 'S', 737.485,  '09:52:43', 'B', 737.4976, '09:54:24',  -2.80),
    ('多', 'B', 735.9898, '09:59:51', 'S', 736.1894, '10:10:44',  18.42),
    ('多', 'B', 739.37,   '11:17:50', 'S', 739.3032, '11:21:19',  -8.22),
    ('空', 'S', 738.69,   '11:27:00', 'B', 739.175,  '11:33:19', -50.04),
    ('空', 'S', 737.9302, '11:47:08', 'B', 738.0692, '11:48:42', -15.44),
    ('多', 'B', 738.60,   '11:49:44', 'S', 738.901,  '11:50:07',  28.56),
    ('多', 'B', 738.19,   '12:29:32', 'S', 737.77,   '12:30:17', -43.54),
]

wb_ta = openpyxl.load_workbook('trades_all.xlsx')
ws_ta = wb_ta.active
start_row = ws_ta.max_row + 1
assert start_row == 803, f"Expected start_row=803, got {start_row}"

for p in pairs:
    pt, oa, op, ot, ca, cp, ct, pnl = p
    ws_ta.append([DATE_STR, ot, 'SPY', op, None, None, 100, oa, 0, DATE_STR, ct])
    ws_ta.append([DATE_STR, ct, 'SPY', cp, pt, pnl, 100, ca, 0, DATE_STR, ot])

wb_ta.save('trades_all.xlsx')
print(f"trades_all.xlsx: appended 20 rows at {start_row}-{start_row+19}")

# ============================================================
# 2. CS交易紀錄.xlsx - 個別識別法 sheet
# ============================================================
wb_cs = openpyxl.load_workbook('CS交易紀錄.xlsx')
ws_cs = wb_cs['個別識別法']
cs_max = ws_cs.max_row
assert cs_max == 1556, f"Expected max_row=1556, got {cs_max}"

ws_cs.append([None, None, None])  # 1557
ws_cs.append([None, None, None])  # 1558

hdr_row = 1559
ws_cs.append([DATE_DOT, DAY_HEADER_B, '損益'])

detail_start = 1560

col1_values = [
    73579.52, -73487.00, -73451.00, 73489.47, -73610.50, 73654.47,
    73746.96, -73749.76, -73598.98, 73617.40, -73937.00, 73928.78,
    73867.46, -73917.50, 73791.48, -73806.92, -73860.00, 73888.56,
    -73819.00, 73775.46,
    -15578.00, 15691.50, -15599.00, 15617.50,
]

pnl_map = {
    1:  92.52,  3:  38.47,  5:  43.97,  7:  -2.80,  9:  18.42,
    11: -8.22, 13: -50.04, 15: -15.44, 17:  28.56, 19: -43.54,
    21: 113.50, 23: 18.50,
}

for i, c1 in enumerate(col1_values):
    r = detail_start + i
    pnl = pnl_map.get(i, None)
    b_formula = f'=+B{r-1}+A{r}'
    ws_cs.cell(row=r, column=1, value=c1)
    ws_cs.cell(row=r, column=2, value=b_formula)
    if pnl is not None:
        ws_cs.cell(row=r, column=3, value=pnl)

detail_end = detail_start + 23  # 1583
spcx_detail_start = detail_start + 20  # 1580
spcx_detail_end = detail_end  # 1583

gokei_row = detail_end + 1  # 1584
ws_cs.cell(row=gokei_row, column=2, value='合計 ')
ws_cs.cell(row=gokei_row, column=3, value=f'=SUM(C{detail_start}:C{detail_end})')

kaken_row = gokei_row + 1  # 1585
ws_cs.cell(row=kaken_row, column=2, value='核驗')
ws_cs.cell(row=kaken_row, column=3, value=f'=+B{detail_end}-B{hdr_row}')

blank1_row = kaken_row + 1  # 1586

spy_ds_start = blank1_row + 1  # 1587
spcx_ds_start = spy_ds_start + 5  # 1592

# SPY summary
ws_cs.cell(row=spy_ds_start, column=1, value=f'{DATE_DOT}當沖')
ws_cs.cell(row=spy_ds_start, column=2, value='買')
ws_cs.cell(row=spy_ds_start, column=3, value=f'=SUMIF(A{detail_start}:A{detail_end},"<0")-C{spcx_ds_start}')
ws_cs.cell(row=spy_ds_start, column=3).number_format = '0.00_);[Red](0.00)'

ws_cs.cell(row=spy_ds_start+1, column=1, value=f'="SPY " & (COUNT(C{detail_start}:C{detail_end})-COUNT(C{spcx_detail_start}:C{spcx_detail_end})) * 100 & "股"')
ws_cs.cell(row=spy_ds_start+1, column=2, value='賣')
ws_cs.cell(row=spy_ds_start+1, column=3, value=f'=SUMIF(A{detail_start}:A{detail_end},">0")-C{spcx_ds_start+1}')
ws_cs.cell(row=spy_ds_start+1, column=3).number_format = '0.00_);[Red](0.00)'

ws_cs.cell(row=spy_ds_start+2, column=2, value='當日損益')
ws_cs.cell(row=spy_ds_start+2, column=3, value=f'=+C{spy_ds_start}+C{spy_ds_start+1}')
ws_cs.cell(row=spy_ds_start+2, column=3).number_format = '0.00_);[Red](0.00)'

ws_cs.cell(row=spy_ds_start+3, column=2, value='當日損益 %')
ws_cs.cell(row=spy_ds_start+3, column=3, value=f'=C{spy_ds_start+2}/ABS(C{spy_ds_start})')
ws_cs.cell(row=spy_ds_start+3, column=3).number_format = '0.00%'

# SPCX summary
ws_cs.cell(row=spcx_ds_start, column=1, value=f'{DATE_DOT} 當沖')
ws_cs.cell(row=spcx_ds_start, column=2, value='買')
ws_cs.cell(row=spcx_ds_start, column=3, value=f'=SUMIF(A{spcx_detail_start}:A{spcx_detail_end},"<0")')
ws_cs.cell(row=spcx_ds_start, column=3).number_format = '0.00_);[Red](0.00)'

ws_cs.cell(row=spcx_ds_start+1, column=1, value=f'="SPCX " & COUNT(C{spcx_detail_start}:C{spcx_detail_end}) * 100 & "股"')
ws_cs.cell(row=spcx_ds_start+1, column=2, value='賣')
ws_cs.cell(row=spcx_ds_start+1, column=3, value=f'=SUMIF(A{spcx_detail_start}:A{spcx_detail_end},">0")')
ws_cs.cell(row=spcx_ds_start+1, column=3).number_format = '0.00_);[Red](0.00)'

ws_cs.cell(row=spcx_ds_start+2, column=2, value='當日損益')
ws_cs.cell(row=spcx_ds_start+2, column=3, value=f'=+C{spcx_ds_start}+C{spcx_ds_start+1}')
ws_cs.cell(row=spcx_ds_start+2, column=3).number_format = '0.00_);[Red](0.00)'

ws_cs.cell(row=spcx_ds_start+3, column=2, value='當日損益 %')
ws_cs.cell(row=spcx_ds_start+3, column=3, value=f'=C{spcx_ds_start+2}/ABS(C{spcx_ds_start})')
ws_cs.cell(row=spcx_ds_start+3, column=3).number_format = '0.00%'

# ============================================================
# 3. cumulative PnL check sheet
# ============================================================
ws_ck = wb_cs['累積損益check']
ck_row = ws_ck.max_row + 1
assert ck_row == 33, f"Expected ck_row=33, got {ck_row}"
ws_ck.cell(row=ck_row, column=1, value=datetime(2026, 6, 24))
ws_ck.cell(row=ck_row, column=2, value=233.90)
ws_ck.cell(row=ck_row, column=3, value=f'=C{ck_row-1}+B{ck_row}')

wb_cs.save('CS交易紀錄.xlsx')
print(f"CS: wrote 6/24 block rows {hdr_row}-{spcx_ds_start+3}")
print(f"cumulative check: added row {ck_row}")

# ============================================================
# 4. Verification
# ============================================================
wb_v = openpyxl.load_workbook('trades_all.xlsx')
ws_v = wb_v.active
print(f"\ntrades_all max_row: {ws_v.max_row}")
for r in range(803, 823):
    vals = [ws_v.cell(r, c).value for c in range(1, 12)]
    print(f"  Row {r}: {vals}")

wb_cv = openpyxl.load_workbook('CS交易紀錄.xlsx')
ws_cv = wb_cv['個別識別法']
print(f"\nCS max_row: {ws_cv.max_row}")
print(f"Header ({hdr_row}): {[ws_cv.cell(hdr_row, c).value for c in range(1,4)]}")
print(f"Detail first ({detail_start}): {[ws_cv.cell(detail_start, c).value for c in range(1,4)]}")
print(f"Detail last ({detail_end}): {[ws_cv.cell(detail_end, c).value for c in range(1,4)]}")
print(f"Gokei ({gokei_row}): {[ws_cv.cell(gokei_row, c).value for c in range(1,4)]}")
print(f"Kaken ({kaken_row}): {[ws_cv.cell(kaken_row, c).value for c in range(1,4)]}")
for r in range(spy_ds_start, spy_ds_start+4):
    print(f"SPY DS row {r}: {[ws_cv.cell(r, c).value for c in range(1,4)]}")
for r in range(spcx_ds_start, spcx_ds_start+4):
    print(f"SPCX DS row {r}: {[ws_cv.cell(r, c).value for c in range(1,4)]}")

ws_ckv = wb_cv['累積損益check']
print(f"\nCumulative row {ck_row}: {[ws_ckv.cell(ck_row, c).value for c in range(1,4)]}")
print("\nAll writes complete.")
