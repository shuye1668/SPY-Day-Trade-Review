import openpyxl
import sys
sys.stdout.reconfigure(encoding='utf-8')

# 1. Revert trades_all.xlsx — delete rows 594-605
print("=== Reverting trades_all.xlsx ===")
wb = openpyxl.load_workbook(r'C:\TradeReview\trades_all.xlsx')
ws = wb.active
print(f"  Current max_row: {ws.max_row}")
if ws.max_row >= 594:
    rows_to_delete = ws.max_row - 593
    ws.delete_rows(594, rows_to_delete)
    print(f"  Deleted {rows_to_delete} rows (594-{593+rows_to_delete})")
    print(f"  New max_row: {ws.max_row}")
wb.save(r'C:\TradeReview\trades_all.xlsx')
print("  Saved.")

# 2. Revert CS交易紀錄.xlsx
print("\n=== Reverting CS交易紀錄.xlsx ===")
wb2 = openpyxl.load_workbook(r'C:\TradeReview\CS交易紀錄.xlsx')
ws2 = wb2['個別識別法']
print(f"  Current max_row: {ws2.max_row}")

# Find and delete the 2026.06.04 block
block_start = None
for r in range(1, ws2.max_row + 1):
    if ws2.cell(r, 1).value == '2026.06.04' and ws2.cell(r, 3).value == '損益':
        block_start = r
        break

if block_start:
    # Delete from block_start to max_row (including trailing blanks)
    rows_to_delete = ws2.max_row - block_start + 1
    # But also delete the 2 blank rows before the header
    start_delete = block_start - 2 if block_start > 2 and ws2.cell(block_start-1, 1).value is None else block_start
    rows_to_delete = ws2.max_row - start_delete + 1
    ws2.delete_rows(start_delete, rows_to_delete)
    print(f"  Deleted {rows_to_delete} rows ({start_delete}-{start_delete+rows_to_delete-1})")
    print(f"  New max_row: {ws2.max_row}")

# Revert 累積損益check — delete last row if it's 2026-06-04
ws3 = wb2['累積損益check']
from datetime import datetime
last_date = ws3.cell(ws3.max_row, 1).value
if isinstance(last_date, datetime) and last_date.year == 2026 and last_date.month == 6 and last_date.day == 4:
    ws3.delete_rows(ws3.max_row, 1)
    print(f"  Deleted 累積損益check row for 2026-06-04")
    print(f"  New max_row: {ws3.max_row}")

wb2.save(r'C:\TradeReview\CS交易紀錄.xlsx')
print("  Saved.")
print("\nRevert complete. Ready to re-run process_0604.py")
