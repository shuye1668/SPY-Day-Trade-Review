import openpyxl
from datetime import datetime

DATE = '2026-06-26'
DATE_DOT = '2026.06.26'

# ===== trades_all.xlsx =====
trades = [
    [DATE, '09:45:51', 'SPY', 728.0175, None, None, 100, 'S', '0', DATE, '09:57:21'],
    [DATE, '09:55:17', 'SPY', 729.585, None, None, 100, 'S', '0', DATE, '09:56:32'],
    [DATE, '09:56:32', 'SPY', 730.18, '空', -61.02, 100, 'B', '0', DATE, '09:55:17'],
    [DATE, '09:57:21', 'SPY', 730.835, '空', -283.27, 100, 'B', '0', DATE, '09:45:51'],
    [DATE, '10:10:54', 'SPY', 732.8499, None, None, 100, 'B', '0', DATE, '10:12:14'],
    [DATE, '10:12:14', 'SPY', 732.88, '多', 1.48, 100, 'S', '0', DATE, '10:10:54'],
    [DATE, '10:21:24', 'SPY', 734.0225, None, None, 100, 'B', '0', DATE, '10:27:48'],
    [DATE, '10:27:48', 'SPY', 733.875, '多', -16.28, 100, 'S', '0', DATE, '10:21:24'],
    [DATE, '10:35:49', 'SPY', 732.98, None, None, 100, 'S', '0', DATE, '10:38:08'],
    [DATE, '10:38:08', 'SPY', 732.49, '空', 47.47, 100, 'B', '0', DATE, '10:35:49'],
    [DATE, '10:40:14', 'SPY', 731.755, None, None, 100, 'S', '0', DATE, '10:41:14'],
    [DATE, '10:41:14', 'SPY', 732.2892, '空', -54.95, 100, 'B', '0', DATE, '10:40:14'],
    [DATE, '10:42:01', 'SPY', 732.4999, None, None, 100, 'B', '0', DATE, '10:46:28'],
    [DATE, '10:46:28', 'SPY', 732.8301, '多', 31.49, 100, 'S', '0', DATE, '10:42:01'],
    [DATE, '10:52:14', 'SPY', 732.93, None, None, 100, 'B', '0', DATE, '11:28:08'],
    [DATE, '10:53:08', 'SPY', 732.53, None, None, 100, 'B', '0', DATE, '10:55:12'],
    [DATE, '10:55:12', 'SPY', 732.545, '多', -0.03, 100, 'S', '0', DATE, '10:53:08'],
    [DATE, '11:05:01', 'SPY', 732.845, None, None, 100, 'B', '0', DATE, '11:23:34'],
    [DATE, '11:16:04', 'SPY', 731.76, None, None, 100, 'B', '0', DATE, '11:17:35'],
    [DATE, '11:17:35', 'SPY', 732.785, '多', 100.97, 100, 'S', '0', DATE, '11:16:04'],
    [DATE, '11:23:34', 'SPY', 733.80, '多', 93.97, 100, 'S', '0', DATE, '11:05:01'],
    [DATE, '11:28:08', 'SPY', 733.8301, '多', 88.48, 100, 'S', '0', DATE, '10:52:14'],
]

wb_ta = openpyxl.load_workbook('trades_all.xlsx')
ws_ta = wb_ta.active
start_row = ws_ta.max_row + 1
for i, t in enumerate(trades):
    r = start_row + i
    for j, v in enumerate(t):
        ws_ta.cell(row=r, column=j+1, value=v)
wb_ta.save('trades_all.xlsx')
print(f'trades_all.xlsx: wrote {len(trades)} rows starting at row {start_row}')

# ===== CS交易紀錄.xlsx =====
wb_cs = openpyxl.load_workbook('CS交易紀錄.xlsx')
ws = wb_cs['個別識別法']

DAY_HEADER_B = 38457.03
BROKER_BAL_OPEN = 38021.78
BROKER_BAL_CLOSE = 37970.09

offset_prev = 37374.26 - 36939.01  # = 435.25
candidate = BROKER_BAL_OPEN + offset_prev
print(f'offset_prev = {offset_prev:.2f}')
print(f'Day Header B candidate = {candidate:.2f}')
assert abs(candidate - DAY_HEADER_B) < 0.01

details = [
    (72800.23, None),
    (72956.98, None),
    (-73018.00, -61.02),
    (-73083.50, -283.27),
    (-73284.99, None),
    (73286.47, 1.48),
    (-73402.25, None),
    (73385.97, -16.28),
    (73296.47, None),
    (-73249.00, 47.47),
    (73173.97, None),
    (-73228.92, -54.95),
    (-73249.99, None),
    (73281.48, 31.49),
    (-73293.00, None),
    (-73253.00, None),
    (73252.97, -0.03),
    (-73284.50, None),
    (-73176.00, None),
    (73276.97, 100.97),
    (73378.47, 93.97),
    (73381.48, 88.48),
]

col1_sum = sum(d[0] for d in details)
pnl_sum = sum(d[1] for d in details if d[1] is not None)
print(f'Sum col1 = {col1_sum:.2f}')
print(f'Sum PnL = {pnl_sum:.2f}')
assert abs(col1_sum - (-51.69)) < 0.01
assert abs(pnl_sum - (-51.69)) < 0.01

header_row = ws.max_row + 3
print(f'CS block starts at row {header_row}')

ws.cell(row=header_row, column=1, value=DATE_DOT)
ws.cell(row=header_row, column=2, value=DAY_HEADER_B)
ws.cell(row=header_row, column=3, value='損益')

detail_start = header_row + 1
for i, (col1, pnl) in enumerate(details):
    r = detail_start + i
    ws.cell(row=r, column=1, value=col1)
    if i == 0:
        ws.cell(row=r, column=2, value=f'=+B{header_row}+A{r}')
    else:
        ws.cell(row=r, column=2, value=f'=+B{r-1}+A{r}')
    if pnl is not None:
        ws.cell(row=r, column=3, value=pnl)

detail_end = detail_start + len(details) - 1

totals_row = detail_end + 1
ws.cell(row=totals_row, column=2, value='合計 ')
ws.cell(row=totals_row, column=3, value=f'=SUM(C{detail_start}:C{detail_end})')

verify_row = totals_row + 1
ws.cell(row=verify_row, column=2, value='核驗')
ws.cell(row=verify_row, column=3, value=f'=+B{detail_end}-B{header_row}')

blank1 = verify_row + 1

spy_buy_row = blank1 + 1
spy_sell_row = spy_buy_row + 1
spy_pnl_row = spy_sell_row + 1
spy_pct_row = spy_pnl_row + 1

ws.cell(row=spy_buy_row, column=1, value=f'{DATE_DOT}當沖')
ws.cell(row=spy_buy_row, column=2, value='買')
ws.cell(row=spy_buy_row, column=3, value=f'=SUMIF(A{detail_start}:A{detail_end},"<0")')

ws.cell(row=spy_sell_row, column=1, value=f'="SPY " & COUNT(C{detail_start}:C{detail_end}) * 100 & "股"')
ws.cell(row=spy_sell_row, column=2, value='賣')
ws.cell(row=spy_sell_row, column=3, value=f'=SUMIF(A{detail_start}:A{detail_end},">0")')

ws.cell(row=spy_pnl_row, column=2, value='當日損益')
ws.cell(row=spy_pnl_row, column=3, value=f'=+C{spy_buy_row}+C{spy_sell_row}')

ws.cell(row=spy_pct_row, column=2, value='當日損益 %')
ws.cell(row=spy_pct_row, column=3, value=f'=C{spy_pnl_row}/ABS(C{spy_buy_row})')

fmt_money = '0.00_);[Red](0.00)'
fmt_pct = '0.00%'
ws.cell(row=spy_buy_row, column=3).number_format = fmt_money
ws.cell(row=spy_sell_row, column=3).number_format = fmt_money
ws.cell(row=spy_pnl_row, column=3).number_format = fmt_money
ws.cell(row=spy_pct_row, column=3).number_format = fmt_pct

# 累積損益check
ws2 = wb_cs['累積損益check']
check_row = ws2.max_row + 1
ws2.cell(row=check_row, column=1, value=datetime(2026, 6, 26))
ws2.cell(row=check_row, column=2, value=-51.69)
ws2.cell(row=check_row, column=3, value=f'=C{check_row-1}+B{check_row}')

print(f'Detail: {detail_start}-{detail_end}, Totals: {totals_row}, Verify: {verify_row}')
print(f'SPY summary: {spy_buy_row}-{spy_pct_row}')
print(f'Cumulative check row: {check_row}')

# Head-tail double offset verification
tail_cs_b = DAY_HEADER_B + col1_sum
head_offset = DAY_HEADER_B - BROKER_BAL_OPEN
tail_offset = tail_cs_b - BROKER_BAL_CLOSE
print(f'head_offset={head_offset:.2f}, tail_offset={tail_offset:.2f}, offset_prev={offset_prev:.2f}')
assert abs(head_offset - offset_prev) < 0.01
assert abs(tail_offset - offset_prev) < 0.01
assert abs(head_offset - tail_offset) < 0.01
print('Head-tail double offset: PASSED')

cash_change = BROKER_BAL_CLOSE - BROKER_BAL_OPEN
assert abs(col1_sum - cash_change) < 0.01
print(f'Independent self-consistency: PASSED (col1_sum={col1_sum:.2f} == cash_change={cash_change:.2f})')

wb_cs.save('CS交易紀錄.xlsx')
print('CS交易紀錄.xlsx saved')
print('\nALL DONE')
