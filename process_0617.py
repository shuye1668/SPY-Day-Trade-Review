import openpyxl
from datetime import datetime

# =====================================================================
# 6/17 Trade Data (from Account Trade History, all SPY, all day trades)
# Format: (edt_time, action, price, col1, pair_idx, pnl_or_none)
# col1: BOT = -(qty*price), SOLD = (qty*price) - 1.57
# Misc Fees estimated at $1.57 per 100-share SOLD
# =====================================================================

trades = [
    ('09:33:01', 'B', 750.4806, -75048.06, 1, None),
    ('09:41:13', 'S', 750.745,   75072.93, 0, 24.87),
    ('09:48:02', 'B', 751.36,   -75136.00, 3, None),
    ('09:48:19', 'S', 751.49,    75147.43, 2, 11.43),
    ('09:55:30', 'B', 751.1218, -75112.18, 5, None),
    ('10:00:04', 'S', 749.7601,  74974.44, 4, -137.74),
    ('10:02:57', 'B', 749.9399, -74993.99, 7, None),
    ('10:03:41', 'S', 750.2071,  75019.14, 6, 25.15),
    ('10:08:18', 'B', 750.96,   -75096.00, 9, None),
    ('10:08:53', 'S', 750.93,    75091.43, 8, -4.57),
    ('10:12:02', 'B', 750.80,   -75080.00, 11, None),
    ('10:12:26', 'S', 750.925,   75090.93, 10, 10.93),
    ('10:24:05', 'B', 751.715,  -75171.50, 13, None),
    ('10:24:36', 'S', 752.0426,  75202.69, 12, 31.19),
    ('10:25:49', 'B', 751.81,   -75181.00, 15, None),
    ('10:26:47', 'S', 751.8101,  75179.44, 14, -1.56),
    ('11:35:20', 'B', 750.569,  -75056.90, 17, None),
    ('11:36:15', 'S', 750.567,   75055.13, 16, -1.77),
    ('11:37:27', 'B', 750.419,  -75041.90, 19, None),
    ('11:57:27', 'S', 750.745,   75072.93, 18, 31.03),
    ('12:21:14', 'B', 750.861,  -75086.10, 21, None),
    ('12:22:43', 'S', 750.94,    75092.43, 20, 6.33),
    ('12:28:40', 'B', 750.883,  -75088.30, 23, None),
    ('12:44:21', 'S', 750.2751,  75025.94, 22, -62.36),
    ('12:50:30', 'B', 749.9899, -74998.99, 25, None),
    ('12:51:36', 'S', 750.25,    75023.43, 24, 24.44),
    ('12:53:24', 'B', 750.005,  -75000.50, 27, None),
    ('12:55:35', 'S', 750.175,   75015.93, 26, 15.43),
    ('12:58:31', 'B', 750.005,  -75000.50, 29, None),
    ('13:11:29', 'S', 749.475,   74945.93, 28, -54.57),
]

# Verify PnL calculations
total_pnl = 0
winners = 0
losers = 0
for i, t in enumerate(trades):
    edt, action, price, col1, pair_idx, pnl = t
    if pnl is not None:
        entry = trades[pair_idx]
        expected_pnl = round(entry[3] + col1, 2)
        if abs(expected_pnl - pnl) > 0.01:
            print(f"PnL MISMATCH at trade {i}: expected {expected_pnl}, got {pnl}")
        total_pnl += pnl
        if pnl > 0:
            winners += 1
        else:
            losers += 1

total_pnl = round(total_pnl, 2)
print(f"Total PnL: {total_pnl}")
print(f"Winners: {winners}, Losers: {losers}, Win rate: {winners}/{winners+losers} = {winners/(winners+losers)*100:.0f}%")

# =====================================================================
# STEP 1: Write to trades_all.xlsx
# =====================================================================
wb = openpyxl.load_workbook("trades_all.xlsx")
ws = wb.active
start_row = ws.max_row + 1
print(f"trades_all: appending from row {start_row}")

date_str = "2026-06-17"
for i, t in enumerate(trades):
    edt, action, price, col1, pair_idx, pnl = t
    row = start_row + i
    pair_trade = trades[pair_idx]

    ws.cell(row, 1, date_str)
    ws.cell(row, 2, edt)
    ws.cell(row, 3, "SPY")
    ws.cell(row, 4, price)

    if pnl is not None:
        type_val = "多"  # 多
        pnl_str = f"+{pnl:.2f}" if pnl >= 0 else f"{pnl:.2f}"
        ws.cell(row, 5, type_val)
        ws.cell(row, 6, pnl_str)
    else:
        ws.cell(row, 5, None)
        ws.cell(row, 6, None)

    ws.cell(row, 7, 100)
    ws.cell(row, 8, action)
    ws.cell(row, 9, 0)
    ws.cell(row, 10, date_str)
    ws.cell(row, 11, pair_trade[0])

wb.save("trades_all.xlsx")
print(f"trades_all.xlsx: wrote {len(trades)} rows (rows {start_row}-{start_row + len(trades) - 1})")
wb.close()

# =====================================================================
# STEP 2: Write to CS交易紀錄.xlsx
# =====================================================================
wb = openpyxl.load_workbook("CS交易紀錄.xlsx")
ws = wb["個別識別法"]  # 個別識別法

# Idempotency check
found_existing = False
for r in range(1, ws.max_row + 1):
    a = ws.cell(r, 1).value
    c = ws.cell(r, 3).value
    if a == "2026.06.17" and c is not None and "損益" in str(c):
        print(f"CS: 6/17 Day Header already exists at row {r}, skipping")
        found_existing = True
        break

if not found_existing:
    cs_start = ws.max_row + 2
    print(f"CS: writing 6/17 block from row {cs_start}")

    day_header_b = -263944.56

    # [1] Day Header
    header_row = cs_start
    ws.cell(header_row, 1, "2026.06.17")
    ws.cell(header_row, 2, day_header_b)
    ws.cell(header_row, 3, "損益")  # 損益

    # [2] Detail rows
    detail_start = header_row + 1

    for i, t in enumerate(trades):
        edt, action, price, col1_val, pair_idx, pnl = t
        r = detail_start + i

        ws.cell(r, 1, round(col1_val, 2))

        if i == 0:
            ws.cell(r, 2, f"=+B{header_row}+A{r}")
        else:
            ws.cell(r, 2, f"=+B{r-1}+A{r}")

        if pnl is not None:
            ws.cell(r, 3, round(pnl, 2))
        else:
            ws.cell(r, 3, None)

    detail_end = detail_start + len(trades) - 1

    # [3] 合計
    sum_row = detail_end + 1
    ws.cell(sum_row, 1, None)
    ws.cell(sum_row, 2, "合計 ")  # 合計 (trailing space)
    ws.cell(sum_row, 3, f"=SUM(C{detail_start}:C{detail_end})")

    # [4] 核驗 (no 留倉/平倉 to subtract)
    check_row = sum_row + 1
    ws.cell(check_row, 1, None)
    ws.cell(check_row, 2, "核驗")  # 核驗
    ws.cell(check_row, 3, f"=+B{detail_end}-B{header_row}")

    # [5] blank
    blank1 = check_row + 1

    # [6] SPY 當沖 4 rows
    spy_buy_row = blank1 + 1
    spy_sell_row = spy_buy_row + 1
    spy_pnl_row = spy_sell_row + 1
    spy_pct_row = spy_pnl_row + 1

    ws.cell(spy_buy_row, 1, "2026.06.17當沖")  # 當沖
    ws.cell(spy_buy_row, 2, "買")  # 買
    ws.cell(spy_buy_row, 3, f'=SUMIF(A{detail_start}:A{detail_end},"<0")')

    ws.cell(spy_sell_row, 1, f'="SPY " & COUNT(C{detail_start}:C{detail_end}) * 100 & "股"')
    ws.cell(spy_sell_row, 2, "賣")  # 賣
    ws.cell(spy_sell_row, 3, f'=SUMIF(A{detail_start}:A{detail_end},">0")')

    ws.cell(spy_pnl_row, 1, None)
    ws.cell(spy_pnl_row, 2, "當日損益")  # 當日損益
    ws.cell(spy_pnl_row, 3, f"=+C{spy_buy_row}+C{spy_sell_row}")

    ws.cell(spy_pct_row, 1, None)
    ws.cell(spy_pct_row, 2, "當日損益 %")  # 當日損益 %
    ws.cell(spy_pct_row, 3, f"=C{spy_pnl_row}/ABS(C{spy_buy_row})")

    # Number formats
    fmt_money = '0.00_);[Red](0.00)'
    fmt_pct = '0.00%'
    ws.cell(spy_buy_row, 3).number_format = fmt_money
    ws.cell(spy_sell_row, 3).number_format = fmt_money
    ws.cell(spy_pnl_row, 3).number_format = fmt_money
    ws.cell(spy_pct_row, 3).number_format = fmt_pct

    print(f"CS: Day Header row {header_row}, Detail {detail_start}-{detail_end}")
    print(f"CS: Sum row {sum_row}, Check row {check_row}")
    print(f"CS: SPY daytrade rows {spy_buy_row}-{spy_pct_row}")

    # =========================================================
    # Update 累積損益check sheet
    # =========================================================
    ws2 = wb["累積損益check"]  # 累積損益check

    found_date = None
    for r2 in range(2, ws2.max_row + 1):
        d = ws2.cell(r2, 1).value
        if d is not None:
            if isinstance(d, datetime) and d.year == 2026 and d.month == 6 and d.day == 17:
                found_date = r2
                break

    if found_date:
        r2 = found_date
        print(f"cum check: updating existing row {r2}")
    else:
        r2 = ws2.max_row + 1
        print(f"cum check: appending at row {r2}")

    ws2.cell(r2, 1, datetime(2026, 6, 17))
    ws2.cell(r2, 2, total_pnl)
    if r2 > 2:
        ws2.cell(r2, 3, f"=C{r2-1}+B{r2}")
    else:
        ws2.cell(r2, 3, total_pnl)

    print(f"cum check: row {r2} = 2026-06-17, PnL={total_pnl}")

wb.save("CS交易紀錄.xlsx")
print("CS交易紀錄.xlsx saved successfully")
wb.close()

print()
print("=== DONE ===")
