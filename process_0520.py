"""
Process US trading day 2026-05-20 (Wednesday)
- 5 round trips, all SPY day trades (Status=0), all wins
- 5/15 holds (200 shares) carried forward unchanged
"""
import openpyxl
from datetime import datetime

# ============================================================
# 1. trades_all.xlsx
# ============================================================
print("=== trades_all.xlsx ===")
wb_trades = openpyxl.load_workbook(r"C:\TradeReview\trades_all.xlsx")
ws_trades = wb_trades.active
print(f"Current max_row: {ws_trades.max_row}")

# Verify no 5/20 data exists
last_date = ws_trades.cell(ws_trades.max_row, 1).value
print(f"Last row date: {last_date}")

# 5 LIFO-paired round trips for 5/20:
# Pair 1: B@734.678 09:37 -> S@735.155 09:40  PnL=+46.17
# Pair 2: B@736.06  09:52 -> S@737.2301 10:18 PnL=+115.47
# Pair 3: B@735.5499 09:34 -> S@737.5001 10:19 PnL=+193.48
# Pair 4: B@738.71  10:31 -> S@739.3294 10:58 PnL=+60.40
# Pair 5: B@739.275 12:51 -> S@739.47   12:55 PnL=+17.96

date_str = "2026-05-20"
rows_trades = [
    # Date, Exec Time(EDT), Symbol, Price, Type, PnL, Shares, Action, Status, Pair_Date, Pair_Time
    [date_str, "09:34:16", "SPY", 735.5499, None, None, 100, "B", "0", date_str, "10:19:20"],
    [date_str, "09:37:01", "SPY", 734.678,  None, None, 100, "B", "0", date_str, "09:40:24"],
    [date_str, "09:40:24", "SPY", 735.155,  "多", "+46.17",  100, "S", "0", date_str, "09:37:01"],
    [date_str, "09:52:07", "SPY", 736.06,   None, None, 100, "B", "0", date_str, "10:18:14"],
    [date_str, "10:18:14", "SPY", 737.2301, "多", "+115.47", 100, "S", "0", date_str, "09:52:07"],
    [date_str, "10:19:20", "SPY", 737.5001, "多", "+193.48", 100, "S", "0", date_str, "09:34:16"],
    [date_str, "10:31:31", "SPY", 738.71,   None, None, 100, "B", "0", date_str, "10:58:23"],
    [date_str, "10:58:23", "SPY", 739.3294, "多", "+60.40",  100, "S", "0", date_str, "10:31:31"],
    [date_str, "12:51:57", "SPY", 739.275,  None, None, 100, "B", "0", date_str, "12:55:00"],
    [date_str, "12:55:00", "SPY", 739.47,   "多", "+17.96",  100, "S", "0", date_str, "12:51:57"],
]

start_row = ws_trades.max_row + 1
for i, row_data in enumerate(rows_trades):
    r = start_row + i
    for c, val in enumerate(row_data, 1):
        ws_trades.cell(r, c, val)

print(f"Appended rows {start_row}-{start_row + len(rows_trades) - 1}")
wb_trades.save(r"C:\TradeReview\trades_all.xlsx")
wb_trades.close()
print("trades_all.xlsx saved.\n")

# ============================================================
# 2. CS交易紀錄.xlsx
# ============================================================
print("=== CS交易紀錄.xlsx ===")
wb_cs = openpyxl.load_workbook(r"C:\TradeReview\CS交易紀錄.xlsx")
ws = wb_cs["個別識別法"]
print(f"Current max_row: {ws.max_row}")

# Idempotency check: look for existing 2026.05.20 Day Header
found_existing = False
for r in range(1, ws.max_row + 1):
    a = ws.cell(r, 1).value
    c = ws.cell(r, 3).value
    if a == "2026.05.20" and c == "損益":
        print(f"WARNING: 2026.05.20 block already exists at row {r}. Skipping.")
        found_existing = True
        break

if not found_existing:
    # Opening balance: from 5/19 block's last detail row (row 962) col B value
    # B962 = -42580.73 (verified from data_only read)
    opening_balance = -42580.73

    # Start 2 blank rows after current max_row (974)
    block_start = ws.max_row + 3  # row 977

    # === Day Header ===
    hr = block_start  # 977
    ws.cell(hr, 1, "2026.05.20")
    ws.cell(hr, 2, opening_balance)
    ws.cell(hr, 3, "損益")

    # === 16 detail rows (978-993) ===
    details = [
        # (col_a, col_c_pnl)  -- col_c_pnl=None means no PnL on this row
        (-73554.99, None),    # BOT @735.5499  row 978
        (-73467.80, None),    # BOT @734.678   row 979
        (73513.97,  46.17),   # SOLD @735.155  row 980 -> pair 1 PnL
        (-73606.00, None),    # BOT @736.06    row 981
        (22116.43,  None),    # SOLD -30 split 1/4  row 982
        (22116.46,  None),    # SOLD -30 split 2/4  row 983
        (22116.43,  None),    # SOLD -30 split 3/4  row 984
        (7372.15,   115.47),  # SOLD -10 split 4/4  row 985 -> pair 2 PnL
        (73748.47,  193.48),  # SOLD @737.5001 row 986 -> pair 3 PnL
        (-73871.00, None),    # BOT @738.71    row 987
        (14786.30,  None),    # SOLD -20 split 1/4  row 988
        (36965.73,  None),    # SOLD -50 split 2/4  row 989
        (5914.44,   None),    # SOLD -8  split 3/4  row 990
        (16264.93,  60.40),   # SOLD -22 split 4/4  row 991 -> pair 4 PnL
        (-73927.50, None),    # BOT @739.275   row 992
        (73945.46,  17.96),   # SOLD @739.47   row 993 -> pair 5 PnL
    ]

    detail_start = hr + 1  # 978
    for i, (col_a, col_c) in enumerate(details):
        r = detail_start + i
        ws.cell(r, 1, col_a)
        # col B: cumulative balance formula
        if i == 0:
            ws.cell(r, 2, f"=+B{hr}+A{r}")
        else:
            ws.cell(r, 2, f"=+B{r-1}+A{r}")
        # col C: PnL (only on exit rows)
        if col_c is not None:
            ws.cell(r, 3, col_c)

    detail_end = detail_start + len(details) - 1  # 993

    # === 合計 row ===
    sum_row = detail_end + 1  # 994
    ws.cell(sum_row, 2, "合計 ")  # trailing space per convention
    ws.cell(sum_row, 3, f"=SUM(C{detail_start}:C{detail_end})")

    # === 核驗 row ===
    check_row = sum_row + 1  # 995
    ws.cell(check_row, 2, "核驗")
    ws.cell(check_row, 3, f"=+B{detail_end}-B{hr}")

    # === blank ===
    blank1 = check_row + 1  # 996

    # === SPY 當沖總結 (4 rows) ===
    ds_row = blank1 + 1  # 997
    ws.cell(ds_row, 1, "2026.05.20當沖")
    ws.cell(ds_row, 2, "買")
    ws.cell(ds_row, 3, f'=SUMIF(A{detail_start}:A{detail_end},"<0")')

    ds_row2 = ds_row + 1  # 998
    ws.cell(ds_row2, 1, f'="SPY " & COUNT(C{detail_start}:C{detail_end}) * 100 & "股"')
    ws.cell(ds_row2, 2, "賣")
    ws.cell(ds_row2, 3, f'=SUMIF(A{detail_start}:A{detail_end},">0")')

    ds_row3 = ds_row2 + 1  # 999
    ws.cell(ds_row3, 2, "當日損益")
    ws.cell(ds_row3, 3, f"=+C{ds_row}+C{ds_row2}")

    ds_row4 = ds_row3 + 1  # 1000
    ws.cell(ds_row4, 2, "當日損益 %")
    ws.cell(ds_row4, 3, f"=C{ds_row3}/ABS(C{ds_row})")

    # === 2 blanks ===
    # rows 1001, 1002

    # === 留倉 section ===
    hold_header_row = ds_row4 + 3  # 1003
    ws.cell(hold_header_row, 1, "累積留倉   200股")

    hold1_row = hold_header_row + 1  # 1004
    ws.cell(hold1_row, 1, "2026.05.15 留倉")
    ws.cell(hold1_row, 2, "買 100股")
    ws.cell(hold1_row, 3, "=C949")  # references 5/18 block hold (correct chain)

    hold2_row = hold1_row + 1  # 1005
    ws.cell(hold2_row, 1, "2026.05.15 留倉")
    ws.cell(hold2_row, 2, "買 100股")
    ws.cell(hold2_row, 3, "=C950")  # references 5/18 block hold (correct chain)

    print(f"CS block written: rows {hr}-{hold2_row}")
    print(f"  Day Header: row {hr}")
    print(f"  Details: rows {detail_start}-{detail_end} (16 rows)")
    print(f"  Sum: row {sum_row}, Check: row {check_row}")
    print(f"  SPY daytrade summary: rows {ds_row}-{ds_row4}")
    print(f"  Holds: rows {hold_header_row}-{hold2_row}")

    # ============================================================
    # 3. 累積損益check sheet
    # ============================================================
    ws2 = wb_cs["累積損益check"]
    cum_row = ws2.max_row + 1  # 12
    ws2.cell(cum_row, 1, datetime(2026, 5, 20))
    ws2.cell(cum_row, 2, 433.48)  # day PnL = 46.17+115.47+193.48+60.40+17.96
    ws2.cell(cum_row, 3, f"=C{cum_row-1}+B{cum_row}")
    print(f"\ncumcheck: row {cum_row} added (2026-05-20, +433.48)")

    wb_cs.save(r"C:\TradeReview\CS交易紀錄.xlsx")
    print("CS交易紀錄.xlsx saved.")

wb_cs.close()
print("\nDone.")
