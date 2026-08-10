#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
cs_from_trades.py — 讀 trades_all.xlsx → append CS交易紀錄.xlsx（2026-07-24 新架構）

【定位】這是 CS交易紀錄 的**唯一正式寫入工具**，純確定性、不需要 AI。
流程分成兩段，職責切乾淨：
    第一段（可能需要 AI）：對帳單 CSV / Gmail草稿 / 截圖  →  trades_all.xlsx
                          （AI 用 PROMPT_轉trades_all.md；有 CSV 時可用 spy_daytrade_engine.py 全自動）
    第二段（本檔，永遠不需要 AI）：trades_all.xlsx  →  CS交易紀錄.xlsx
AI 之後的角色只剩「檢查 / 修正 / 維護」。AI 不能用時，人工照《操作手冊》跑本檔即可。

【為什麼只靠 trades_all 就夠】trades_all 沒有手續費欄，但淨額可**精確反推**（非估算）：
    買進(B)無手續費 → net = -(price × shares)              ← 精確
    賣出(S)         → net = pnl - buy_net                   ← 由該組配對的損益反推
反推出的手續費會逐筆還原真實值（如 1.54 / 1.55 / 1.56 各不同），
符合 CLAUDE.md §1「手續費一律取實際值、嚴禁估算」。已對 2026-07-22/23 共 36 筆驗證逐筆相同。

【唯一需要人工提供的數字】當日期初餘額（券商對帳單 BAL 行「Cash balance at the start of business day」）。
trades_all 內不含此資訊，故用 --open-bal 傳入。強烈建議同時給 --close-bal 以啟用頭尾雙錨點檢核。

用法：
    python cs_from_trades.py --date 2026-07-22 --open-bal 37874.49 --close-bal 37668.99
    python cs_from_trades.py --date 2026-07-22 --open-bal 37874.49 --close-bal 37668.99 --commit

預設 DRY-RUN（不碰任何檔案）。--commit 才寫入。

安全（對齊 CLAUDE.md）：
  §0 任何不確定 → 中止不寫，回報需人工確認。絕不靜默寫估算/推算/猜值。
  §2 寫入位置用值掃描 last_nonempty_row，嚴禁 ws.max_row。
  §3 寫入前自動備份。 §4 寫入後全核驗，失敗自動還原。 §5 檢查 ~$ 鎖檔。
  冪等：同一天重複跑不會重複寫。
"""
import argparse
import json
import os
import shutil
import sys
from datetime import datetime

import openpyxl

_BASE = os.path.dirname(os.path.abspath(__file__))
TRADES_PATH = os.path.join(_BASE, "trades_all.xlsx")
CS_PATH = os.path.join(_BASE, "CS交易紀錄.xlsx")
OFFSET_STATE = os.path.join(_BASE, "offset_state.json")
SHEET_DETAIL = "個別識別法"

FMT_NUM = r"0.00_);[Red]\(0.00\)"
FMT_PCT = "0.00%"
FEE_MAX = 3.00          # 每 100 股手續費合理上限（超過視為異常）
EPS = 0.011             # 金額比對容差


def balances_from_statement(csv_path, date):
    """從對帳單 CSV 取出指定 EDT 交易日的 (期初餘額, 收盤餘額, session 列數)。

    刻意直接沿用引擎的 parse_statement/group_sessions，而不是自己再寫一套
    CSV 解析 —— 期初餘額的定義（BAL 列、CST→EDT 換日、跨午夜歸屬哪一個
    session）全部由引擎決定。自己重寫一份遲早會與 writer 的認定分岔，
    而那正是最難察覺、後果最嚴重的一類錯誤。

    找不到該日 session 時回傳 (None, None, 0)。
    """
    sys.path.insert(0, _BASE)
    from spy_daytrade_engine import parse_statement, group_sessions

    with open(csv_path, encoding="utf-8-sig") as fh:
        sessions = group_sessions(parse_statement(fh.read()))
    for s in sessions:
        # s.trades 要 process_session() 之後才有內容；這裡只取餘額，
        # 所以回報原始成交列數 s.fills 才是真實的。
        if s.edt_date == date:
            return s.open_bal, s.close_bal, len(s.fills)
    return None, None, 0


# ── 共用小工具 ───────────────────────────────────────────────────────────────
def last_nonempty_row(ws, ncols=5):
    for r in range(ws.max_row, 0, -1):
        if any(ws.cell(r, c).value not in (None, "") for c in range(1, ncols + 1)):
            return r
    return 0


def check_lock(path):
    d, base = os.path.split(path)
    if os.path.exists(os.path.join(d, "~$" + base)):
        raise RuntimeError(f"§5 鎖檔存在（Excel 開著？）：{path} — 請關閉 Excel 後重跑")


def load_offset():
    if not os.path.exists(OFFSET_STATE):
        raise RuntimeError(f"找不到 {OFFSET_STATE}")
    with open(OFFSET_STATE, encoding="utf-8") as f:
        return float(json.load(f)["offset"])


def _norm_date(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()[:10]


def _f(v):
    return None if v in (None, "") else float(v)


# ── 1. 讀 trades_all，推導淨額 ───────────────────────────────────────────────
def load_day(date, trades_path=TRADES_PATH):
    """回傳 (rows, problems)。rows 已按時間排序，每筆含 net_cash / pnl。"""
    ws = openpyxl.load_workbook(trades_path, data_only=True).active
    all_rows = []
    for r in range(2, ws.max_row + 1):
        d = _norm_date(ws.cell(r, 1).value)
        if d != date:
            continue
        all_rows.append({
            "row": r, "date": d,
            "time": str(ws.cell(r, 2).value).strip(),
            "symbol": str(ws.cell(r, 3).value or "").strip(),
            "price": _f(ws.cell(r, 4).value),
            "type": ws.cell(r, 5).value,
            "pnl": _f(ws.cell(r, 6).value),
            "shares": int(_f(ws.cell(r, 7).value) or 0),
            "action": str(ws.cell(r, 8).value or "").strip().upper(),
            "status": str(ws.cell(r, 9).value or "").strip(),
            "pair_time": str(ws.cell(r, 11).value or "").strip(),
            "net": None,
        })
    problems = []
    if not all_rows:
        problems.append(f"trades_all 中找不到 {date} 的任何交易")
        return all_rows, problems

    for x in all_rows:
        if x["status"] != "0":
            problems.append(
                f"{x['time']} Status={x['status']}（留倉/跨日平倉）— 非當沖，需人工處理（見 CLAUDE.md）")
        if x["shares"] != 100:
            problems.append(f"{x['time']} 股數={x['shares']}（非 100）— 需人工確認")
        if x["action"] not in ("B", "S"):
            problems.append(f"{x['time']} Action={x['action']!r} 無效")
    if problems:
        return all_rows, problems

    all_rows.sort(key=lambda x: x["time"])
    # 配對：有 pnl 的是出場列，pair_time 指向進場列
    used = set()
    pairs = []
    for ex in all_rows:
        if ex["pnl"] is None:
            continue
        cand = [y for y in all_rows
                if y["time"] == ex["pair_time"] and y["pnl"] is None
                and id(y) not in used and y["action"] != ex["action"]]
        if not cand:
            problems.append(f"{ex['time']} 找不到配對的進場列（Pair_Time={ex['pair_time']}）")
            continue
        en = cand[0]
        used.add(id(en))
        pairs.append((en, ex))

    unpaired = [x for x in all_rows if x["pnl"] is None and id(x) not in used]
    for x in unpaired:
        problems.append(f"{x['time']} 進場列沒有對應的出場列（未平倉？）— 需人工確認")
    if problems:
        return all_rows, problems

    # 推導淨額：B 精確、S 由 pnl 反推
    for en, ex in pairs:
        b, s = (en, ex) if en["action"] == "B" else (ex, en)
        if b["action"] != "B" or s["action"] != "S":
            problems.append(f"{ex['time']} 配對方向異常（應一買一賣）")
            continue
        b_net = round(-b["price"] * b["shares"], 2)
        s_net = round(ex["pnl"] - b_net, 2)
        fee = round(round(s["price"] * s["shares"], 2) - s_net, 2)
        if not (-0.001 <= fee <= FEE_MAX):
            problems.append(
                f"{s['time']} 反推手續費 {fee:.2f} 超出合理範圍 0~{FEE_MAX} — 資料可能有誤")
        b["net"], s["net"] = b_net, s_net

    if any(x["net"] is None for x in all_rows):
        problems.append("有交易未能推導出淨額")
    return all_rows, problems


# ── 2. 檢核 ─────────────────────────────────────────────────────────────────
def validate(rows, open_bal, close_bal, offset):
    """回傳 (problems, info)。close_bal 可為 None（則跳過頭尾/自洽檢核並警告）。"""
    problems, info = [], {}
    s_net = round(sum(x["net"] for x in rows), 2)
    s_pnl = round(sum(x["pnl"] for x in rows if x["pnl"] is not None), 2)
    info["Σcol1"], info["Σpnl"] = s_net, s_pnl
    if abs(s_net - s_pnl) > EPS:
        problems.append(f"§4.1 Σcol1({s_net}) ≠ Σpnl({s_pnl})，差 {s_net - s_pnl:+.2f}")

    header_B = round(open_bal + offset, 2)
    chain_close = round(header_B + s_net, 2)
    info["Day Header B"], info["chain 收盤"] = header_B, chain_close
    if close_bal is not None:
        if abs(s_net - round(close_bal - open_bal, 2)) > EPS:
            problems.append(
                f"§4.3 Σcol1({s_net}) ≠ 券商收-開({round(close_bal - open_bal, 2)})"
                f" — 可能漏筆或手續費有誤")
        head = round(header_B - open_bal, 2)
        tail = round(chain_close - close_bal, 2)
        info["head_offset"], info["tail_offset"] = head, tail
        if abs(head - offset) > EPS or abs(tail - offset) > EPS:
            problems.append(
                f"§4.2 頭尾 offset 不符：head={head} tail={tail} 期望={offset}")
    else:
        info["警告"] = "未提供 --close-bal，略過頭尾雙錨點與自洽檢核（強烈建議補上）"
    return problems, info


# ── 3. 寫入 ─────────────────────────────────────────────────────────────────
def _detect_tail_hold(ws, last):
    """檔尾若是『累積留倉』區段，回傳 (起始列, 內容)；否則 (None, [])。"""
    start = None
    for r in range(last, 0, -1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a.startswith("累積留倉"):
            start = r
            break
        if isinstance(a, str) and ws.cell(r, 3).value == "損益":
            break
    if start is None:
        return None, []
    return start, [(ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value)
                   for r in range(start, last + 1)]


def day_exists(ws, dot):
    for r in range(1, last_nonempty_row(ws) + 1):
        if ws.cell(r, 1).value == dot and ws.cell(r, 3).value == "損益":
            return True
    return False


def write_block(ws, date, rows, header_B):
    dot = date.replace("-", ".")
    last = last_nonempty_row(ws)
    hold_start, hold_rows = _detect_tail_hold(ws, last)
    if hold_start is not None:
        ws.delete_rows(hold_start, last - hold_start + 1)
        while last_nonempty_row(ws) < ws.max_row:
            ws.delete_rows(ws.max_row, 1)

    hr = last_nonempty_row(ws) + 3           # §2：前留 2 空行
    ws.cell(hr, 1, dot); ws.cell(hr, 2, header_B); ws.cell(hr, 3, "損益")
    d0 = hr + 1
    for i, x in enumerate(rows):
        r = d0 + i
        ws.cell(r, 1, x["net"])
        ws.cell(r, 2, f"=+B{hr}+A{r}" if i == 0 else f"=+B{r-1}+A{r}")
        if x["pnl"] is not None:
            ws.cell(r, 3, x["pnl"])
    dN = d0 + len(rows) - 1

    tot = dN + 1
    ws.cell(tot, 2, "合計 "); ws.cell(tot, 3, f"=SUM(C{d0}:C{dN})")
    ws.cell(tot, 3).number_format = FMT_NUM
    chk = tot + 1
    ws.cell(chk, 2, "核驗"); ws.cell(chk, 3, f"=+B{dN}-B{hr}")
    ws.cell(chk, 3).number_format = FMT_NUM

    buy = chk + 2                             # 中間空 1 行
    ws.cell(buy, 1, f"{dot}當沖"); ws.cell(buy, 2, "買")
    ws.cell(buy, 3, f'=SUMIF(A{d0}:A{dN},"<0")'); ws.cell(buy, 3).number_format = FMT_NUM
    sell = buy + 1
    ws.cell(sell, 1, f'="SPY " & COUNT(C{d0}:C{dN}) * 100 & "股"'); ws.cell(sell, 2, "賣")
    ws.cell(sell, 3, f'=SUMIF(A{d0}:A{dN},">0")'); ws.cell(sell, 3).number_format = FMT_NUM
    pnl = sell + 1
    ws.cell(pnl, 2, "當日損益"); ws.cell(pnl, 3, f"=+C{buy}+C{sell}")
    ws.cell(pnl, 3).number_format = FMT_NUM
    pct = pnl + 1
    ws.cell(pct, 2, "當日損益 %"); ws.cell(pct, 3, f"=C{pnl}/ABS(C{buy})")
    ws.cell(pct, 3).number_format = FMT_PCT

    if hold_rows:
        ws.append([None, None, None])
        for a, b, c in hold_rows:
            ws.append([a, b, c])
    return hr, d0, dN


def write_cumulative(wb, date, day_pnl):
    cs = wb.worksheets[1]                     # 累積損益check
    last = last_nonempty_row(cs, ncols=3)
    r = last + 1
    y, m, d = [int(v) for v in date.split("-")]
    cs.cell(r, 1, datetime(y, m, d))
    cs.cell(r, 2, round(day_pnl, 2))
    cs.cell(r, 3, f"=C{last}+B{r}")


def post_verify(ws, hr, d0, dN, open_bal, close_bal, offset):
    errs = []
    col1 = [ws.cell(r, 1).value for r in range(d0, dN + 1)]
    colC = [ws.cell(r, 3).value for r in range(d0, dN + 1)
            if isinstance(ws.cell(r, 3).value, (int, float))]
    s1, sC = round(sum(col1), 2), round(sum(colC), 2)
    hb = ws.cell(hr, 2).value
    if abs(s1 - sC) > EPS:
        errs.append(f"§4.1 Σcol1({s1})≠Σpnl({sC})")
    if close_bal is not None:
        if abs(round(hb - open_bal, 2) - offset) > EPS or \
           abs(round(hb + s1 - close_bal, 2) - offset) > EPS:
            errs.append("§4.2 頭尾 offset 不符")
        if abs(s1 - round(close_bal - open_bal, 2)) > EPS:
            errs.append("§4.3 Σcol1 ≠ 收-開")
    blanks = 0
    for r in range(hr - 1, 0, -1):
        if all(ws.cell(r, c).value in (None, "") for c in range(1, 6)):
            blanks += 1
        else:
            break
    if blanks != 2:
        errs.append(f"§4.4 header 前空白列={blanks}≠2")
    for i, r in enumerate(range(d0, dN + 1)):
        want = f"=+B{hr}+A{r}" if i == 0 else f"=+B{r-1}+A{r}"
        if ws.cell(r, 2).value != want:
            errs.append(f"§4.5 B{r} 公式不符")
            break
    return errs, sC


# ── main ────────────────────────────────────────────────────────────────────
def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    ap = argparse.ArgumentParser(
        description="讀 trades_all.xlsx，append 當日區塊到 CS交易紀錄.xlsx")
    ap.add_argument("--date", required=True, help="交易日 YYYY-MM-DD")
    ap.add_argument("--open-bal", type=float, default=None,
                    help="券商當日期初餘額（對帳單 BAL 行）；未給則需 --from-statement")
    ap.add_argument("--close-bal", type=float, default=None,
                    help="券商當日收盤餘額（最後一筆 BALANCE）；強烈建議提供")
    ap.add_argument("--from-statement", metavar="CSV",
                    help="從對帳單 CSV 自動帶出期初/收盤餘額（省略路徑則用 "
                         "_inbox\\<date>.csv）。Boss PC 同步過來的原始對帳單就在那")
    ap.add_argument("--trades", default=TRADES_PATH)
    ap.add_argument("--cs", default=CS_PATH)
    ap.add_argument("--commit", action="store_true", help="實際寫入（預設 dry-run）")
    a = ap.parse_args()

    # ── 自動帶入期初/收盤餘額 ────────────────────────────────────────────
    # 期初餘額（BAL 行）過去只能靠人工從對帳單/OCR 找，是每天最容易出錯也最
    # 花時間的一步。Boss PC 會把原始對帳單 CSV 一起同步過來，直接讀就好。
    if a.from_statement is not None or a.open_bal is None:
        stmt = a.from_statement or os.path.join(_BASE, "_inbox", f"{a.date}.csv")
        if not os.path.exists(stmt):
            print(f"🔴 找不到對帳單：{stmt}")
            print("   請改用 --open-bal/--close-bal 手動指定，或確認 git pull 已同步。")
            sys.exit(1)
        try:
            ob, cb, nrow = balances_from_statement(stmt, a.date)
        except Exception as e:
            print(f"🔴 對帳單解析失敗（{stmt}）：{type(e).__name__}: {e}")
            sys.exit(1)
        if ob is None:
            print(f"🔴 對帳單裡找不到 {a.date} 的 BAL 列（期初餘額）。")
            print("   OCR 模式常漏抓 BAL/DOI/JRN 列 —— 這正是 offset_state.json")
            print("   記載 2026-08-05 需要回推的原因。請補齊對帳單後重跑。")
            sys.exit(1)
        if a.open_bal is None:
            a.open_bal = ob
        if a.close_bal is None:
            a.close_bal = cb
        print(f"[自動帶入] 來源 {os.path.basename(stmt)}（對帳單 {nrow} 筆成交）"
              f" 期初={a.open_bal} 收盤={a.close_bal}")

    offset = load_offset()
    rows, problems = load_day(a.date, a.trades)
    if not problems:
        p2, info = validate(rows, a.open_bal, a.close_bal, offset)
        problems += p2
    else:
        info = {}

    print(f"=== {a.date} ===  trades_all 取得 {len(rows)} 筆")
    for x in rows:
        print(f"  {x['time']}  {x['action']}  {x['price']:>10.4f}  "
              f"net={x['net'] if x['net'] is not None else '?':>12}  "
              f"{(x['type'] or '-'):<3} pnl={'' if x['pnl'] is None else format(x['pnl'], '.2f')}")
    for k, v in info.items():
        print(f"  {k}: {v}")

    if problems:
        print(f"\n🔴 {len(problems)} 處需人工確認，不寫入：")
        for p in problems:
            print(f"   • {p}")
        print("\n請依《操作手冊》第 4 節處理，或交負責人。切勿手動改數字硬湊。")
        sys.exit(1)

    print("\n🟢 全部檢核通過")
    if not a.commit:
        print("[DRY-RUN] 未寫入。確認無誤後加 --commit 實際寫入。")
        return

    check_lock(a.cs)
    bkp = a.cs[:-5] + f"_backup_{datetime.now():%Y%m%d}.xlsx"
    shutil.copy2(a.cs, bkp)
    print(f"§3 已備份：{os.path.basename(bkp)}")
    try:
        wb = openpyxl.load_workbook(a.cs)
        ws = wb[SHEET_DETAIL]
        dot = a.date.replace("-", ".")
        if day_exists(ws, dot):
            print(f"— {a.date} 已存在於 CS（冪等跳過，未重複寫入）")
            return
        header_B = round(a.open_bal + offset, 2)
        hr, d0, dN = write_block(ws, a.date, rows, header_B)
        errs, day_pnl = post_verify(ws, hr, d0, dN, a.open_bal, a.close_bal, offset)
        if errs:
            raise RuntimeError("；".join(errs))
        write_cumulative(wb, a.date, day_pnl)
        wb.save(a.cs)
        print(f"✅ 寫入完成：Day Header 第 {hr} 列，{dN-d0+1} 筆明細，"
              f"當日損益 {day_pnl:+.2f}，並通過 §4 全核驗")
    except Exception as e:
        shutil.copy2(bkp, a.cs)
        print(f"❌ 寫入失敗，已還原備份：{e}")
        raise


if __name__ == "__main__":
    main()
