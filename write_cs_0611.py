import openpyxl

wb = openpyxl.load_workbook('CS交易紀錄.xlsx')
ws = wb['個別識別法']

# ============================================================
# Idempotency check: look for existing 2026.06.11 Day Header
# ============================================================
for r in range(1, ws.max_row + 1):
    a = ws.cell(r, 1).value
    c = ws.cell(r, 3).value
    if a == '2026.06.11' and c == '損益':
        print(f'CS交易紀錄已含 2026.06.11 區塊（第 {r} 列），不重寫')
        exit()

# ============================================================
# 期初餘額 chain: B(6/11) = B(6/10) + 6/10 day trade PnL
# ============================================================
# 6/10 Day Header at row 1273, B = -41683.01
# 6/10 day trade PnL = +195.65
b_6_11 = -41683.01 + 195.65  # = -41487.36
print(f'B(6/11) = {b_6_11:.2f}')

# ============================================================
# Detail rows: all 19 merged trades for 6/11 in EDT order
# ============================================================
# (col1, col_c_pnl) - col_c is None for entries, PnL for day trade exits
detail = [
    (-72778.50, None),     # 1: BOT 100 @727.785 09:36:33
    (72919.48, 140.98),    # 2: SOLD 100 @729.21 09:49:23
    (-73051.00, None),     # 3: BOT 100 @730.51 09:57:29
    (73090.47, 39.47),     # 4: SOLD 100 @730.92 10:01:18
    (-72703.00, None),     # 5: BOT 100 @727.03 10:20:01
    (72853.48, 150.48),    # 6: SOLD 100 @728.55 10:22:58
    (-72909.00, None),     # 7: BOT 100 @729.09 10:30:13
    (72865.48, -43.52),    # 8: SOLD 100 @728.67 10:35:39
    (-72985.90, None),     # 9: BOT 100 @729.859 10:37:06
    (72965.57, -20.33),    # 10: SOLD 100 @729.6709 10:42:57
    (-72643.48, None),     # 11: BOT 100 @726.4348 10:56:03
    (72690.88, 47.40),     # 12: SOLD 100 @726.924 11:07:55
    (-72702.99, None),     # 13: BOT 100 @727.0299 11:10:51
    (72709.98, 6.99),      # 14: SOLD 100 @727.115 11:11:50
    (-72707.99, None),     # 15: BOT 100 @727.0799 11:12:27
    (72749.48, 41.49),     # 16: SOLD 100 @727.51 11:13:08
    (-72794.59, None),     # 17: BOT 100 @727.9459 12:33:06 (留倉 entry)
    (-72853.50, None),     # 18: BOT 100 @728.535 12:43:51 (day trade entry)
    (72981.48, 127.98),    # 19: SOLD 100 @729.83 13:29:03 (day trade exit)
]

# ============================================================
# Write block
# ============================================================
current_row = ws.max_row + 1  # 1315

# 2 blank rows
blank1 = current_row
blank2 = current_row + 1
current_row += 2  # now 1317

# [1] Day Header
header_row = current_row
ws.cell(header_row, 1, '2026.06.11')
ws.cell(header_row, 2, b_6_11)
ws.cell(header_row, 3, '損益')
print(f'Day Header at row {header_row}')
current_row += 1

# [2] Detail section
detail_start = current_row
hold_entry_row = None  # track the 留倉 entry row for SUMIF adjustment

for i, (col1, pnl) in enumerate(detail):
    r = current_row
    ws.cell(r, 1, col1)
    if i == 0:
        ws.cell(r, 2, f'=+B{header_row}+A{r}')
    else:
        ws.cell(r, 2, f'=+B{r-1}+A{r}')
    if pnl is not None:
        ws.cell(r, 3, pnl)

    # Track 留倉 entry (index 16, the BOT @727.9459)
    if i == 16:
        hold_entry_row = r

    current_row += 1

detail_end = current_row - 1
print(f'Detail rows {detail_start}-{detail_end} ({detail_end - detail_start + 1} rows)')
print(f'Hold entry at row {hold_entry_row}')

# [3] 合計
sum_row = current_row
ws.cell(sum_row, 1, None)
ws.cell(sum_row, 2, '合計 ')
ws.cell(sum_row, 3, f'=SUM(C{detail_start}:C{detail_end})')
print(f'合計 at row {sum_row}')
current_row += 1

# [4] 核驗
verify_row = current_row
ws.cell(verify_row, 1, None)
ws.cell(verify_row, 2, '核驗')
ws.cell(verify_row, 3, f'=+B{detail_end}-B{header_row}')
print(f'核驗 at row {verify_row}')
current_row += 1

# [5] blank
current_row += 1

# [6] SPY 當沖總結 4 rows
# Detail has 1 留倉 entry (row hold_entry_row), need to subtract from 買 SUMIF
buy_row = current_row
ws.cell(buy_row, 1, '2026.06.11當沖')
ws.cell(buy_row, 2, '買')
ws.cell(buy_row, 3, f'=SUMIF(A{detail_start}:A{detail_end},"<0")-A{hold_entry_row}')
print(f'當沖買 at row {buy_row}: SUMIF minus hold entry A{hold_entry_row}')
current_row += 1

sell_row = current_row
ws.cell(sell_row, 1, f'="SPY " & COUNT(C{detail_start}:C{detail_end}) * 100 & "股"')
ws.cell(sell_row, 2, '賣')
ws.cell(sell_row, 3, f'=SUMIF(A{detail_start}:A{detail_end},">0")')
current_row += 1

pnl_row = current_row
ws.cell(pnl_row, 1, None)
ws.cell(pnl_row, 2, '當日損益')
ws.cell(pnl_row, 3, f'=+C{buy_row}+C{sell_row}')
current_row += 1

pct_row = current_row
ws.cell(pct_row, 1, None)
ws.cell(pct_row, 2, '當日損益 %')
ws.cell(pct_row, 3, f'=C{pnl_row}/ABS(C{buy_row})')
current_row += 1

# [7] blank
current_row += 1

# [11] 留倉區 - 1 new hold (BOT @727.9459) + 3 prior holds still open
# Total holds: 4 (rows 616/617/636 from prior + new 727.9459)
hold_header_row = current_row
ws.cell(hold_header_row, 1, '累積留倉   400股')
ws.cell(hold_header_row, 2, None)
ws.cell(hold_header_row, 3, None)
current_row += 1

# Prior holds: reference their CS交易紀錄 留倉 rows
# From CS交易紀錄 row 1312: '2026.06.05 留倉' 買 100股 =C1268
# From CS交易紀錄 row 1313: '2026.06.05 留倉' 買 100股 =C1269
# From CS交易紀錄 row 1314: '2026.06.09 留倉' 買 100股 =C1270
# These are the PRIOR day's 留倉 rows. For 6/11, I need to reference them.
# But actually, per §5.5.10: "當日新留倉區" only lists NEW holds from today.
# Prior holds that are still open don't get re-listed (they were listed in their originating day).
# Wait, let me re-read...

# §5.5.10: "當日 LIFO 配對後仍有未平倉部位(Status=1)"
# "('累積留倉   {N}股', None, None)"
# "每筆一列"

# Looking at existing data: row 1311 says '累積留倉   300股' and rows 1312-1314 list all 3 holds.
# So the 留倉區 lists ALL current holds (including prior ones).

# But §6.3 says: "當日新增留倉須寫入總結訊息; 舊有留倉不列"
# That's for the summary message only. The CS交易紀錄 留倉區 lists ALL holds.

# For CS交易紀錄, list all 4 holds:
# 6/5 hold 1: reference 6/10's row 1312 col C (which = C1268)
ws.cell(current_row, 1, '2026.06.05 留倉')
ws.cell(current_row, 2, '買 100股')
ws.cell(current_row, 3, f'=C1312')
current_row += 1

# 6/5 hold 2: reference 6/10's row 1313 col C
ws.cell(current_row, 1, '2026.06.05 留倉')
ws.cell(current_row, 2, '買 100股')
ws.cell(current_row, 3, f'=C1313')
current_row += 1

# 6/9 hold: reference 6/10's row 1314 col C
ws.cell(current_row, 1, '2026.06.09 留倉')
ws.cell(current_row, 2, '買 100股')
ws.cell(current_row, 3, f'=C1314')
current_row += 1

# New hold: BOT 100 @727.9459 - col1 = -72794.59 (hardcoded)
ws.cell(current_row, 1, '2026.06.11 留倉')
ws.cell(current_row, 2, '買 100股')
ws.cell(current_row, 3, -72794.59)
current_row += 1

print(f'留倉區 rows {hold_header_row}-{current_row-1}')

# ============================================================
# 累積損益check sheet
# ============================================================
ws2 = wb['累積損益check']
max_r2 = ws2.max_row
print(f'\n累積損益check max_row: {max_r2}')

# Check last row
last_date = ws2.cell(max_r2, 1).value
last_pnl = ws2.cell(max_r2, 2).value
last_cum = ws2.cell(max_r2, 3).value
print(f'Last row {max_r2}: date={last_date} pnl={last_pnl} cum={last_cum}')

# Add 6/11 row
from datetime import datetime
new_row = max_r2 + 1
ws2.cell(new_row, 1, datetime(2026, 6, 11))
ws2.cell(new_row, 2, 490.94)  # day trade PnL only (no 平倉 today)
ws2.cell(new_row, 3, f'=C{max_r2}+B{new_row}')
print(f'Added 累積損益check row {new_row}: date=2026-06-11 pnl=490.94')

# Save
wb.save('CS交易紀錄.xlsx')
print('\nCS交易紀錄.xlsx saved successfully')
